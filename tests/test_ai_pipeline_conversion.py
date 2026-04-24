"""Tests for AI pipeline Stage 3 — conversion."""

from __future__ import annotations

import sys
import zipfile
from io import BytesIO
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from automation.ai_pipeline.conversion import (
    ConvertedArtifact,
    ProjectConversion,
    convert_project,
    load_conversion,
    save_conversion,
)


# ---------------------------------------------------------------------------
# Fixtures — reuse the helpers from typology tests via direct copy
# (keeps tests independent; tiny enough to duplicate)
# ---------------------------------------------------------------------------


def _noisy_png_bytes(size: int = 256) -> bytes:
    import random
    from PIL import Image

    random.seed(42)
    img = Image.new("RGB", (size, size))
    pixels = [(random.randint(0, 255), random.randint(0, 255), random.randint(0, 255))
              for _ in range(size * size)]
    img.putdata(pixels)
    buf = BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _make_pdf_text(text: str, with_image: bool = False) -> bytes:
    import fitz
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), text)
    if with_image:
        page.insert_image(fitz.Rect(100, 300, 356, 556), stream=_noisy_png_bytes(256))
    data = doc.tobytes()
    doc.close()
    return data


def _make_pdf_scanned() -> bytes:
    import fitz
    doc = fitz.open()
    page = doc.new_page()
    page.insert_image(fitz.Rect(50, 50, 500, 500), stream=_noisy_png_bytes(400))
    data = doc.tobytes()
    doc.close()
    return data


def _make_xlsx(sheets: dict[str, list[list]]) -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_docx(text: str) -> bytes:
    # Real DOCX via python-docx — pandoc needs a proper file
    import docx
    d = docx.Document()
    d.add_heading("Title", level=1)
    d.add_paragraph(text)
    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()


def _make_project(tmp_path: Path) -> Path:
    project = tmp_path / "demo"
    project.mkdir()

    (project / "PENETROS.pdf").write_bytes(_make_pdf_text("This is PENETROS page content. " * 10))
    (project / "SONDEIG_scanned.pdf").write_bytes(_make_pdf_scanned())
    (project / "PLAN_with_image.pdf").write_bytes(_make_pdf_text("Plan text. " * 20, with_image=True))
    (project / "BUDGET.xlsx").write_bytes(_make_xlsx({
        "Costs": [["item", "value"], ["Labour", 1000], ["Materials", 500]],
        "Summary": [["total"], [1500]],
    }))
    (project / "letter.docx").write_bytes(_make_docx("Dear Client, please find enclosed…"))
    (project / "notes.txt").write_text("Field notes: soil is granular.")
    # Image at root (treated as image passthrough)
    (project / "foto.png").write_bytes(_noisy_png_bytes(200))
    # Dev-only (skipped, won't get converted)
    (project / "PDF").mkdir()
    (project / "PDF" / "old_report.pdf").write_bytes(_make_pdf_text("old report"))
    return project


# ---------------------------------------------------------------------------
# Shape + counts
# ---------------------------------------------------------------------------


def test_convert_project_returns_model(tmp_path):
    project = _make_project(tmp_path)
    conv = convert_project(project)
    assert isinstance(conv, ProjectConversion)
    assert conv.project_path == str(project.resolve())
    assert conv.converted_at.endswith("+00:00")


def test_skipped_files_are_not_converted(tmp_path):
    """Files with useful=False (e.g. PDF/ dev-only) should not appear as artifacts."""
    project = _make_project(tmp_path)
    conv = convert_project(project)
    source_paths = {a.source_path for a in conv.artifacts}
    assert "PDF/old_report.pdf" not in source_paths


def test_conversion_produces_expected_formats(tmp_path):
    project = _make_project(tmp_path)
    conv = convert_project(project)
    formats = {a.format for a in conv.artifacts if not a.skipped}
    assert "md" in formats        # PDF text, DOCX, Excel not in this simple case
    assert "csv" in formats       # Excel sheets
    assert "png" in formats       # PDF scanned pages
    assert "passthrough" in formats  # image + text


# ---------------------------------------------------------------------------
# Per-strategy converters
# ---------------------------------------------------------------------------


def test_pdf_text_becomes_markdown_per_page(tmp_path):
    project = _make_project(tmp_path)
    conv = convert_project(project)
    pdf_arts = conv.artifacts_of("PENETROS.pdf")
    assert len(pdf_arts) == 1
    a = pdf_arts[0]
    assert a.format == "md"
    assert a.page == 1
    assert a.strategy_used == "pdf_to_markdown"

    md = (project / a.path).read_text(encoding="utf-8")
    assert "PENETROS" in md


def test_pdf_mixed_has_images_manifest(tmp_path):
    project = _make_project(tmp_path)
    conv = convert_project(project)
    arts = conv.artifacts_of("PLAN_with_image.pdf")
    # One page markdown + one images.md
    formats = sorted(a.format for a in arts)
    assert formats.count("md") >= 2
    # The images.md should have a ref to extracted image
    images_md_art = next((a for a in arts if a.path.endswith("images.md")), None)
    assert images_md_art is not None
    content = (project / images_md_art.path).read_text(encoding="utf-8")
    assert "Imatges embegudes" in content
    assert "extracted/" in content


def test_pdf_scanned_becomes_png_per_page(tmp_path):
    project = _make_project(tmp_path)
    conv = convert_project(project)
    arts = conv.artifacts_of("SONDEIG_scanned.pdf")
    assert len(arts) >= 1
    a = arts[0]
    assert a.format == "png"
    assert a.page == 1
    assert a.strategy_used == "pdf_pages_to_images"
    assert (project / a.path).is_file()


def test_docx_becomes_markdown_body(tmp_path):
    project = _make_project(tmp_path)
    conv = convert_project(project)
    arts = conv.artifacts_of("letter.docx")
    body_art = next((a for a in arts if a.path.endswith("body.md")), None)
    assert body_art is not None
    assert body_art.format == "md"
    md = (project / body_art.path).read_text(encoding="utf-8")
    assert "Dear Client" in md


def test_excel_per_sheet_csv(tmp_path):
    project = _make_project(tmp_path)
    conv = convert_project(project)
    arts = conv.artifacts_of("BUDGET.xlsx")
    csv_arts = [a for a in arts if a.format == "csv"]
    # Two sheets → two CSVs
    assert len(csv_arts) == 2
    sheet_names = {a.sheet for a in csv_arts}
    assert sheet_names == {"Costs", "Summary"}
    costs = next(a for a in csv_arts if a.sheet == "Costs")
    data = (project / costs.path).read_text(encoding="utf-8")
    assert "Labour,1000" in data


def test_text_passthrough_references_original(tmp_path):
    project = _make_project(tmp_path)
    conv = convert_project(project)
    arts = conv.artifacts_of("notes.txt")
    assert len(arts) == 1
    a = arts[0]
    assert a.format == "passthrough"
    # passthrough points at the original path, no copy
    assert a.path == "notes.txt"


def test_image_passthrough_references_original(tmp_path):
    project = _make_project(tmp_path)
    conv = convert_project(project)
    arts = conv.artifacts_of("foto.png")
    assert len(arts) == 1
    a = arts[0]
    assert a.format == "passthrough"
    assert a.path == "foto.png"


# ---------------------------------------------------------------------------
# Dedup + re-run
# ---------------------------------------------------------------------------


def test_rerun_is_idempotent(tmp_path):
    project = _make_project(tmp_path)
    conv1 = convert_project(project)
    conv2 = convert_project(project)
    # Same set of artifact paths, same byte sizes
    assert {(a.path, a.bytes_written) for a in conv1.artifacts if not a.skipped} == \
           {(a.path, a.bytes_written) for a in conv2.artifacts if not a.skipped}


def test_rerun_preserves_disk_when_source_unchanged(tmp_path):
    project = _make_project(tmp_path)
    conv1 = convert_project(project)
    md_art = next(a for a in conv1.artifacts if a.format == "md" and a.source_path.endswith("PENETROS.pdf"))
    md_path = project / md_art.path
    mtime_before = md_path.stat().st_mtime

    # Second run — dedup should skip the write
    import time; time.sleep(0.1)  # ensure mtime resolution
    convert_project(project)
    mtime_after = md_path.stat().st_mtime
    assert mtime_after == mtime_before, "dedup should skip re-write for unchanged content"


# ---------------------------------------------------------------------------
# Strategy filter
# ---------------------------------------------------------------------------


def test_strategy_filter_limits_converters(tmp_path):
    project = _make_project(tmp_path)
    conv = convert_project(project, strategies={"pdf_to_markdown"})
    # Only pdf_text source gets converted (PENETROS.pdf)
    sources = {a.source_path for a in conv.artifacts}
    assert sources == {"PENETROS.pdf"}


# ---------------------------------------------------------------------------
# Source chain propagation
# ---------------------------------------------------------------------------


def test_source_chain_is_propagated_to_artifact(tmp_path):
    project = _make_project(tmp_path)
    # Classify first to inject a fake source_chain
    from automation.ai_pipeline.typology import classify_project
    typ = classify_project(project, extract_images=False)
    # Find the PENETROS entry and inject a chain (simulating it came from a .msg)
    for f in typ.files:
        if f.path == "PENETROS.pdf":
            f.source_chain = ["test.msg", "attachment:PENETROS.pdf"]
            break
    conv = convert_project(project, typology=typ)
    art = next(a for a in conv.artifacts if a.source_path == "PENETROS.pdf" and a.format == "md")
    assert "test.msg" in art.source_chain
    assert any(s.startswith("attachment:") for s in art.source_chain)


# ---------------------------------------------------------------------------
# Query helpers
# ---------------------------------------------------------------------------


def test_artifacts_of_returns_correct_group(tmp_path):
    project = _make_project(tmp_path)
    conv = convert_project(project)
    arts = conv.artifacts_of("BUDGET.xlsx")
    assert all(a.source_path == "BUDGET.xlsx" for a in arts)
    assert len(arts) >= 2


def test_artifacts_by_format(tmp_path):
    project = _make_project(tmp_path)
    conv = convert_project(project)
    assert all(a.format == "csv" for a in conv.artifacts_by_format("csv"))
    assert all(a.format == "md" for a in conv.artifacts_by_format("md"))


def test_per_source_manifest(tmp_path):
    project = _make_project(tmp_path)
    conv = convert_project(project)
    assert "BUDGET.xlsx" in conv.per_source
    assert "PENETROS.pdf" in conv.per_source
    # PDF/old_report.pdf is dev-only, should not appear
    assert "PDF/old_report.pdf" not in conv.per_source


# ---------------------------------------------------------------------------
# Eva summary
# ---------------------------------------------------------------------------


def test_eva_summary_is_catalan_and_mentions_counts(tmp_path):
    project = _make_project(tmp_path)
    conv = convert_project(project)
    joined = "\n".join(conv.eva_summary).lower()
    assert "fitxers útils" in joined
    assert "artefactes" in joined


# ---------------------------------------------------------------------------
# Cache
# ---------------------------------------------------------------------------


def test_save_load_roundtrip(tmp_path):
    project = _make_project(tmp_path)
    conv = convert_project(project)
    out = save_conversion(conv, project)
    assert out == project / "validation" / "ai_conversion.json"
    loaded = load_conversion(project)
    assert loaded is not None
    assert len(loaded.artifacts) == len(conv.artifacts)
    assert loaded.total_bytes == conv.total_bytes


def test_load_returns_none_when_absent(tmp_path):
    project = _make_project(tmp_path)
    assert load_conversion(project) is None


# ---------------------------------------------------------------------------
# Errors
# ---------------------------------------------------------------------------


def test_rejects_nonexistent_project(tmp_path):
    with pytest.raises(ValueError):
        convert_project(tmp_path / "no-such-dir")
