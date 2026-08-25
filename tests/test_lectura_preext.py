"""Tests de l'experiment de pre-extracció determinista (`automation/lectura/preext.py`)
i les eines que hi van lligades (`scripts/write_doc_json.py`, `scripts/render_clip.py`).

Vegeu la sonda `docs/wizard-headless/mesures/probes/2026-08-24-tall-stream-json/turns.md`
i el docstring de `automation/lectura/preext.py`.

Fixtures: fitxers SINTÈTICS (fitz/openpyxl/PIL, mai binaris al repo) a `tmp_path`,
tret del `.msg` real (`reference-material/3001621 CASTELLAR DEL VALLES/...`, només
lectura — mai s'hi escriu res) i del PDF de mostra de Bell-lloc per a
`render_clip.py` (`reference-material/4001612 BELL-LLOC/25.0647/A.01.pdf`).
Si algun d'aquests dos fitxers reals no hi és, el test corresponent fa `pytest.skip`.
"""

from __future__ import annotations

import hashlib
import json
import subprocess
import sys
import time
from pathlib import Path

import pytest

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

from automation.lectura.preext import augment_inventory, preext_document  # noqa: E402
from automation.lectura.runner import safe_doc_name  # noqa: E402

WRITE_DOC_JSON = PROJECT_ROOT / "scripts" / "write_doc_json.py"
RENDER_CLIP = PROJECT_ROOT / "scripts" / "render_clip.py"
MSG_FIXTURE = (
    PROJECT_ROOT / "reference-material" / "3001621 CASTELLAR DEL VALLES"
    / "Re_ ESTUDI GEOTÈCNIC VERSIÓ CASTELLAR DEL VALLES.msg"
)
SAMPLE_PDF = PROJECT_ROOT / "reference-material" / "4001612 BELL-LLOC" / "25.0647" / "A.01.pdf"


# ---------------------------------------------------------------------------
# Utilitats de fixture
# ---------------------------------------------------------------------------


def _md5_of(path: Path) -> str:
    return hashlib.md5(path.read_bytes()).hexdigest()


def _entry(rel_path: str, project: Path, route: str = "claude") -> dict:
    return {"path": rel_path, "md5": _md5_of(project / rel_path), "route": route}


def _write_two_page_pdf(path: Path) -> None:
    """Pàgina 1 apaïsada, pàgina 2 vertical, totes dues amb prou text."""
    import fitz

    doc = fitz.open()
    p1 = doc.new_page(width=800, height=400)
    p1.insert_text((72, 200), "Pagina apaisada amb text suficient per superar el llindar de vint caracters.")
    p2 = doc.new_page(width=400, height=800)
    p2.insert_text((36, 400), "Pagina vertical amb text suficient per superar el llindar tambe aqui.")
    doc.save(str(path))
    doc.close()


def _write_n_page_pdf(path: Path, n: int) -> None:
    import fitz

    doc = fitz.open()
    for i in range(n):
        page = doc.new_page(width=400, height=600)
        page.insert_text((36, 300), f"Pagina numero {i + 1} amb text suficient per no ser brossa.")
    doc.save(str(path))
    doc.close()


def _write_xlsx(path: Path) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "fitxa"
    ws1["A1"] = "capçalera"
    ws1["N19"] = 3001621
    ws2 = wb.create_sheet("altre")
    ws2["B2"] = "valor"
    wb.save(str(path))


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def test_pdf_two_pages_clean_text_no_halves(tmp_path):
    """Text net (producer per defecte de fitz, ≥20 alfanumèrics, proporció alta)
    -> `text_ok` true a totes dues pàgines i CAP fitxer de meitats (v2: les
    meitats només es generen per a pàgines amb `text_ok` false)."""
    project = tmp_path / "proj"
    project.mkdir()
    _write_two_page_pdf(project / "doc2p.pdf")
    preext_root = tmp_path / "_preext"

    meta = preext_document(project, _entry("doc2p.pdf", project), preext_root)

    assert meta["kind"] == "pdf"
    assert meta.get("error") is None
    assert meta["pages"] == 2
    assert meta["pages_truncated"] is False
    assert len(meta["page_sizes_pt"]) == 2
    assert "producer" in meta
    assert meta.get("text_suspect_reason") is None

    doc_dir = preext_root / meta["dir"]
    assert "apaisada" in (doc_dir / "page-1.txt").read_text(encoding="utf-8")
    assert (doc_dir / "page-1.png").exists()
    assert not (doc_dir / "page-1-left.png").exists()
    assert not (doc_dir / "page-1-right.png").exists()

    assert "vertical" in (doc_dir / "page-2.txt").read_text(encoding="utf-8")
    assert (doc_dir / "page-2.png").exists()
    assert not (doc_dir / "page-2-top.png").exists()
    assert not (doc_dir / "page-2-bottom.png").exists()

    assert meta["text_ok"]["1"] is True
    assert meta["text_ok"]["2"] is True
    assert meta["halves_pages"] == []

    for rel in meta["files"]:
        assert (preext_root / rel).exists()


def test_pdf_mixed_pages_halves_only_for_garbage_text_page(tmp_path):
    """Pàgina amb text net -> sense meitats; pàgina amb text brossa (pocs
    alfanumèrics / proporció baixa) -> meitats, orientació segons mida."""
    import fitz

    project = tmp_path / "proj"
    project.mkdir()
    path = project / "mixed.pdf"
    doc = fitz.open()
    p1 = doc.new_page(width=400, height=600)
    p1.insert_text((36, 300), "Pagina neta amb text llegible suficient per superar el llindar de vint caracters.")
    p2 = doc.new_page(width=400, height=600)
    p2.insert_text((36, 300), '\t \t !"#$%&\'()' * 10)
    doc.save(str(path))
    doc.close()
    preext_root = tmp_path / "_preext"

    meta = preext_document(project, _entry("mixed.pdf", project), preext_root)

    assert meta.get("error") is None
    assert meta["text_ok"]["1"] is True
    assert meta["text_ok"]["2"] is False
    assert meta["halves_pages"] == [2]

    doc_dir = preext_root / meta["dir"]
    assert not (doc_dir / "page-1-top.png").exists()
    assert not (doc_dir / "page-1-bottom.png").exists()
    assert (doc_dir / "page-2-top.png").exists()
    assert (doc_dir / "page-2-bottom.png").exists()


def test_pdf_distiller_producer_recorded_but_does_not_override_clean_text(tmp_path):
    """Descoberta empírica (Castellar, `ACCEPTACIO/PRESSUPOST GEOTEC...pdf`):
    el productor Acrobat Distiller/PScript5 NO implica per si sol text brossa
    -- l'Eva l'usa tant en annexos amb text brossa (ràtio 0,27) com en
    pressupostos amb text net (ràtio 0,99). `text_suspect_reason` queda
    anotat com a informació, però `text_ok` es decideix pel contingut."""
    import fitz

    project = tmp_path / "proj"
    project.mkdir()
    path = project / "distiller.pdf"
    doc = fitz.open()
    for i in range(2):
        page = doc.new_page(width=400, height=600)
        page.insert_text((36, 300), f"Pagina numero {i + 1} amb text net i suficient per no ser brossa normalment.")
    doc.set_metadata({"producer": "Acrobat Distiller 15.0 (Windows)"})
    doc.save(str(path))
    doc.close()
    preext_root = tmp_path / "_preext"

    meta = preext_document(project, _entry("distiller.pdf", project), preext_root)

    assert meta.get("error") is None
    assert meta["text_ok"]["1"] is True
    assert meta["text_ok"]["2"] is True
    assert "Distiller" in meta["text_suspect_reason"]
    assert meta["halves_pages"] == []

    doc_dir = preext_root / meta["dir"]
    assert not (doc_dir / "page-1-top.png").exists()
    assert not (doc_dir / "page-2-top.png").exists()
    # el .txt s'escriu igualment
    assert (doc_dir / "page-1.txt").exists()


def test_pdf_distiller_producer_with_garbage_text_is_still_flagged_by_content(tmp_path):
    """Rèplica de `ANNEXES/3001621_sondeig.pdf` (Castellar): mateix productor
    Distiller que el test anterior, però amb text brossa real -> `text_ok`
    false (per contingut, no pel productor) i meitats generades."""
    import fitz

    project = tmp_path / "proj"
    project.mkdir()
    path = project / "sondeig.pdf"
    doc = fitz.open()
    page = doc.new_page(width=400, height=600)
    page.insert_text((36, 300), '\t \t !"#$%&\'()' * 10)
    doc.set_metadata({"producer": "Acrobat Distiller 15.0 (Windows)", "creator": "PScript5.dll Version 5.2.2"})
    doc.save(str(path))
    doc.close()
    preext_root = tmp_path / "_preext"

    meta = preext_document(project, _entry("sondeig.pdf", project), preext_root)

    assert meta.get("error") is None
    assert meta["text_ok"]["1"] is False
    assert "Distiller" in meta["text_suspect_reason"]
    assert meta["halves_pages"] == [1]

    doc_dir = preext_root / meta["dir"]
    assert (doc_dir / "page-1-top.png").exists()
    assert (doc_dir / "page-1-bottom.png").exists()


def test_pdf_fourteen_pages_truncated_after_twelve(tmp_path):
    project = tmp_path / "proj"
    project.mkdir()
    _write_n_page_pdf(project / "big.pdf", 14)
    preext_root = tmp_path / "_preext"

    meta = preext_document(project, _entry("big.pdf", project), preext_root)

    assert meta["pages"] == 14
    assert meta["pages_truncated"] is True
    assert len(meta["page_sizes_pt"]) == 14
    assert meta["halves_pages"] == []  # text net a totes: cap meitat calia

    doc_dir = preext_root / meta["dir"]
    assert (doc_dir / "page-12.png").exists()
    assert (doc_dir / "page-13.txt").exists()
    assert (doc_dir / "page-14.txt").exists()
    assert not (doc_dir / "page-13.png").exists()
    assert not (doc_dir / "page-14.png").exists()
    assert not (doc_dir / "page-13-left.png").exists()


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------


def test_xlsx_cells_and_csv(tmp_path):
    project = tmp_path / "proj"
    project.mkdir()
    _write_xlsx(project / "fitxa.xlsx")
    preext_root = tmp_path / "_preext"

    meta = preext_document(project, _entry("fitxa.xlsx", project), preext_root)

    assert meta["kind"] == "excel"
    assert meta.get("error") is None
    assert meta["colors"] is True
    doc_dir = preext_root / meta["dir"]

    cells = (doc_dir / "sheet-0.cells.txt").read_text(encoding="utf-8")
    assert "N19\t3001621" in cells
    assert (doc_dir / "sheet-0.csv").exists()
    assert (doc_dir / "sheet-1.cells.txt").exists()

    # cap cel·la de color: sheet-i.colors.txt existeix però buit
    assert (doc_dir / "sheet-0.colors.txt").read_text(encoding="utf-8") == ""
    assert meta["sheets"][0]["colored_cells"] == 0

    names = {s["name"] for s in meta["sheets"]}
    assert names == {"fitxa", "altre"}


def test_xls_cells_and_csv(tmp_path):
    xlwt = pytest.importorskip("xlwt")

    project = tmp_path / "proj"
    project.mkdir()
    xls_path = project / "fitxa.xls"
    wb = xlwt.Workbook()
    ws = wb.add_sheet("fitxa")
    ws.write(18, 13, 3001621)  # N19 (fila 19, columna N -> 0-based 18,13)
    wb.save(str(xls_path))
    preext_root = tmp_path / "_preext"

    meta = preext_document(project, _entry("fitxa.xls", project), preext_root)

    assert meta["kind"] == "excel"
    assert meta.get("error") is None
    doc_dir = preext_root / meta["dir"]
    cells = (doc_dir / "sheet-0.cells.txt").read_text(encoding="utf-8")
    assert "N19\t3001621" in cells


def test_xlsx_colors_txt_has_only_colored_cells(tmp_path):
    """Cel·la buida amb fons vermell (senyal de nivell freàtic per color, cas
    Alcoletge) + cel·la amb font vermella + cel·la normal -> colors.txt amb
    exactament les dues primeres, valor buit conservat com a camp buit."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    project = tmp_path / "proj"
    project.mkdir()
    path = project / "colors.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "fitxa"
    ws["A1"] = "normal"
    ws["B2"].fill = PatternFill(fill_type="solid", fgColor="FFFF0000")  # buida, fons vermell
    ws["C3"] = "amb font vermella"
    ws["C3"].font = Font(color="FFFF0000")
    wb.save(str(path))
    preext_root = tmp_path / "_preext"

    meta = preext_document(project, _entry("colors.xlsx", project), preext_root)

    assert meta["kind"] == "excel"
    assert meta.get("error") is None
    assert meta["sheets"][0]["colored_cells"] == 2

    doc_dir = preext_root / meta["dir"]
    lines = (doc_dir / "sheet-0.colors.txt").read_text(encoding="utf-8").splitlines()
    assert len(lines) == 2
    by_ref = {ln.split("\t")[0]: ln for ln in lines}

    assert "B2" in by_ref
    assert "bg=#FF0000" in by_ref["B2"]
    assert by_ref["B2"].endswith("\t")  # valor buit conservat com a camp buit

    assert "C3" in by_ref
    assert "font=#FF0000" in by_ref["C3"]
    assert by_ref["C3"].endswith("amb font vermella")


def test_xls_colors_txt_with_pattern_fill(tmp_path):
    xlwt = pytest.importorskip("xlwt")

    project = tmp_path / "proj"
    project.mkdir()
    xls_path = project / "colors.xls"
    wb = xlwt.Workbook()
    ws = wb.add_sheet("fitxa")
    ws.write(0, 0, "normal")
    red_style = xlwt.easyxf("pattern: pattern solid, fore_colour red")
    ws.write(1, 1, "", red_style)  # B2 (0-based fila1, columna1), buida amb fons vermell
    wb.save(str(xls_path))
    preext_root = tmp_path / "_preext"

    meta = preext_document(project, _entry("colors.xls", project), preext_root)

    assert meta["kind"] == "excel"
    assert meta.get("error") is None
    assert meta.get("colors_error") is None
    assert meta["sheets"][0]["colored_cells"] >= 1

    doc_dir = preext_root / meta["dir"]
    lines = (doc_dir / "sheet-0.colors.txt").read_text(encoding="utf-8").splitlines()
    assert any(ln.startswith("B2\t") and "bg=" in ln for ln in lines)


# ---------------------------------------------------------------------------
# Correu (.msg real, només lectura)
# ---------------------------------------------------------------------------


def test_msg_body_and_attachments(tmp_path):
    if not MSG_FIXTURE.exists():
        pytest.skip(".msg de mostra no disponible")

    project = MSG_FIXTURE.parent
    preext_root = tmp_path / "_preext"

    meta = preext_document(project, _entry(MSG_FIXTURE.name, project), preext_root)

    assert meta["kind"] == "msg"
    assert meta.get("error") is None
    doc_dir = preext_root / meta["dir"]

    body = (doc_dir / "body.txt").read_text(encoding="utf-8")
    assert "Subject:" in body
    assert "From:" in body

    assert isinstance(meta["attachments"], list)
    for att in meta["attachments"]:
        assert {"name", "size", "md5", "already_in_folder", "preext_dir"} <= att.keys()

    before = {p.relative_to(project) for p in project.rglob("*")}
    preext_document(project, _entry(MSG_FIXTURE.name, project), preext_root)
    after = {p.relative_to(project) for p in project.rglob("*")}
    assert before == after, "la pre-extracció mai ha d'escriure dins project_path"


# ---------------------------------------------------------------------------
# Imatge
# ---------------------------------------------------------------------------


def test_image_reduced_to_max_long_side(tmp_path):
    from PIL import Image

    project = tmp_path / "proj"
    project.mkdir()
    img_path = project / "photo.png"
    Image.new("RGB", (3000, 2000), color=(120, 130, 140)).save(img_path)
    preext_root = tmp_path / "_preext"

    meta = preext_document(project, _entry("photo.png", project), preext_root)

    assert meta["kind"] == "image"
    assert meta.get("error") is None
    doc_dir = preext_root / meta["dir"]

    with Image.open(doc_dir / "image.png") as out_im:
        assert max(out_im.size) == 1600
    assert meta["original_size"] == img_path.stat().st_size


# ---------------------------------------------------------------------------
# Idempotència
# ---------------------------------------------------------------------------


def test_idempotent_by_md5_then_redo_on_change(tmp_path):
    from PIL import Image

    project = tmp_path / "proj"
    project.mkdir()
    img_path = project / "photo.jpg"
    Image.new("RGB", (200, 100), (1, 2, 3)).save(img_path)
    preext_root = tmp_path / "_preext"

    entry1 = _entry("photo.jpg", project)
    meta1 = preext_document(project, entry1, preext_root)
    meta_path = preext_root / meta1["dir"] / "meta.json"
    mtime1 = meta_path.stat().st_mtime_ns

    meta2 = preext_document(project, entry1, preext_root)
    assert meta_path.stat().st_mtime_ns == mtime1
    assert meta2 == meta1

    # Alguns filesystems (p. ex. 9p/WSL) arrodoneixen mtime al segon: dona marge
    # perquè el "redo" sigui detectable per mtime, no només per contingut.
    time.sleep(1.1)

    Image.new("RGB", (80, 60), (9, 9, 9)).save(img_path)
    entry2 = _entry("photo.jpg", project)
    assert entry2["md5"] != entry1["md5"]

    meta3 = preext_document(project, entry2, preext_root)
    assert meta3["source_md5"] == entry2["md5"]
    assert meta3 != meta1
    assert meta_path.stat().st_mtime_ns != mtime1


# ---------------------------------------------------------------------------
# Errors (mai fan caure la pre-extracció)
# ---------------------------------------------------------------------------


def test_pdf_error_is_captured_in_meta(tmp_path):
    project = tmp_path / "proj"
    project.mkdir()
    bad_pdf = project / "broken.pdf"
    bad_pdf.write_bytes(b"aixo no es un pdf de veritat, nomes bytes de proves")
    preext_root = tmp_path / "_preext"

    meta = preext_document(project, _entry("broken.pdf", project), preext_root)

    assert meta.get("error")
    assert meta["kind"] == "pdf"  # deduit per extensio malgrat l'error
    assert meta["files"] == []


def test_dwg_and_other_extensions_are_unsupported(tmp_path):
    project = tmp_path / "proj"
    project.mkdir()
    (project / "planol.dwg").write_bytes(b"contingut binari fals de dwg")
    preext_root = tmp_path / "_preext"

    meta = preext_document(project, _entry("planol.dwg", project), preext_root)

    assert meta["kind"] == "unsupported"
    assert meta.get("error") is None
    assert meta["unsupported"] is True
    assert meta["files"] == []


# ---------------------------------------------------------------------------
# augment_inventory
# ---------------------------------------------------------------------------


def test_augment_inventory_preext_only_on_claude_entries(tmp_path):
    project = tmp_path / "proj"
    project.mkdir()
    _write_two_page_pdf(project / "annex.pdf")

    out_dir = tmp_path / "lectura"
    out_dir.mkdir()
    inv = {
        "generated": "now", "project": "proj",
        "files": [
            {"path": "annex.pdf", "size": (project / "annex.pdf").stat().st_size, "md5": _md5_of(project / "annex.pdf"),
             "mtime": "now", "route": "claude", "doc_type_hint": "altre", "priority": 1},
            {"path": "skip.jpg", "size": 10, "md5": "deadbeef" * 4, "mtime": "now",
             "route": "skip", "doc_type_hint": "foto", "priority": 0},
            {"path": "pressupost.pdf", "size": 10, "md5": "cafebabe" * 4, "mtime": "now",
             "route": "python", "doc_type_hint": "pressupost_g3", "priority": 0},
        ],
        "duplicates": [],
    }
    inv_path = out_dir / "_inventory.json"
    inv_path.write_text(json.dumps(inv, ensure_ascii=False), encoding="utf-8")

    result = augment_inventory(project, inv_path)

    by_path = {f["path"]: f for f in result["files"]}
    assert "preext" in by_path["annex.pdf"]
    pe = by_path["annex.pdf"]["preext"]
    assert Path(pe["dir"]).is_absolute() and Path(pe["dir"]).exists()
    assert Path(pe["meta"]).exists()
    assert pe["files"]
    for fp in pe["files"]:
        assert Path(fp).is_absolute()
        assert Path(fp).exists()

    assert "preext" not in by_path["skip.jpg"]
    assert "preext" not in by_path["pressupost.pdf"]

    assert result["preext"]["n_docs"] == 1
    assert result["preext"]["n_errors"] == 0
    assert Path(result["preext"]["root"]).is_absolute()

    reloaded = json.loads(inv_path.read_text(encoding="utf-8"))
    assert "preext" in reloaded
    assert "preext" in {f["path"]: f for f in reloaded["files"]}["annex.pdf"]


def test_augment_inventory_error_does_not_raise(tmp_path):
    project = tmp_path / "proj"
    project.mkdir()
    bad_pdf = project / "broken.pdf"
    bad_pdf.write_bytes(b"garbage, not a real pdf")

    out_dir = tmp_path / "lectura"
    out_dir.mkdir()
    inv = {
        "generated": "now", "project": "proj",
        "files": [
            {"path": "broken.pdf", "size": bad_pdf.stat().st_size, "md5": _md5_of(bad_pdf),
             "mtime": "now", "route": "claude", "doc_type_hint": "altre", "priority": 1},
        ],
        "duplicates": [],
    }
    inv_path = out_dir / "_inventory.json"
    inv_path.write_text(json.dumps(inv, ensure_ascii=False), encoding="utf-8")

    result = augment_inventory(project, inv_path)  # no ha de llençar

    assert result["preext"]["n_docs"] == 1
    assert result["preext"]["n_errors"] == 1
    assert result["files"][0]["preext"]["error"]


def test_augment_inventory_idempotent_second_call_zero_reprocess(tmp_path):
    project = tmp_path / "proj"
    project.mkdir()
    _write_two_page_pdf(project / "annex.pdf")

    out_dir = tmp_path / "lectura"
    out_dir.mkdir()
    inv = {
        "generated": "now", "project": "proj",
        "files": [
            {"path": "annex.pdf", "size": (project / "annex.pdf").stat().st_size, "md5": _md5_of(project / "annex.pdf"),
             "mtime": "now", "route": "claude", "doc_type_hint": "altre", "priority": 1},
        ],
        "duplicates": [],
    }
    inv_path = out_dir / "_inventory.json"
    inv_path.write_text(json.dumps(inv, ensure_ascii=False), encoding="utf-8")

    result1 = augment_inventory(project, inv_path)
    meta_path = Path(result1["files"][0]["preext"]["meta"])
    mtime1 = meta_path.stat().st_mtime_ns

    # el proper augment_inventory rellegeix _inventory.json (ja augmentat) del disc.
    result2 = augment_inventory(project, inv_path)
    assert meta_path.stat().st_mtime_ns == mtime1
    assert result2["preext"]["n_errors"] == 0


# ---------------------------------------------------------------------------
# scripts/write_doc_json.py (subprocess)
# ---------------------------------------------------------------------------


def _run_write_doc_json(args: list[str], stdin_text: str) -> subprocess.CompletedProcess:
    return subprocess.run(
        [sys.executable, str(WRITE_DOC_JSON), *args],
        input=stdin_text, capture_output=True, text=True, cwd=str(PROJECT_ROOT),
    )


def test_write_doc_json_valid_payload_writes_canonical_name(tmp_path):
    out_dir = tmp_path / "out"
    payload = {
        "source_path": "25.0647/A.01.pdf", "source_md5": "abc123", "schema_version": 1, "skill_version": "1.0",
        "context": {"authority_for": ["architect_name"]},
        "tier_a": [{"concept_id": "architect_name", "value": "Jordi Bosch", "location": "caixetí",
                    "quote": "Jordi Bosch", "confidence": 0.9}],
        "tables": {"soil_levels": [{"nom": "1er nivell"}], "dpsh_tests": []},
        "not_present": [],
    }
    r = _run_write_doc_json(
        ["--out", str(out_dir), "--expect-source-path", "25.0647/A.01.pdf", "--expect-md5", "abc123"],
        json.dumps(payload),
    )

    assert r.returncode == 0, r.stderr
    assert r.stdout.strip().startswith("OK ")
    assert "tier_a=1" in r.stdout
    assert "soil_levels" in r.stdout
    assert "not_present=0" in r.stdout

    out_path = out_dir / (safe_doc_name("25.0647/A.01.pdf") + ".json")
    assert out_path.exists()
    assert json.loads(out_path.read_text(encoding="utf-8")) == payload
    assert list(out_dir.glob("*.tmp")) == []


def test_write_doc_json_source_path_mismatch_exits_2_writes_nothing(tmp_path):
    out_dir = tmp_path / "out"
    payload = {"source_path": "x.pdf", "source_md5": "m5", "schema_version": 1,
               "context": {"authority_for": []}, "tier_a": [], "not_present": []}

    r = _run_write_doc_json(["--out", str(out_dir), "--expect-source-path", "y.pdf"], json.dumps(payload))

    assert r.returncode == 2
    assert "ERROR" in r.stderr
    assert not out_dir.exists()


def test_write_doc_json_authority_for_without_tier_a_warns_but_succeeds(tmp_path):
    out_dir = tmp_path / "out"
    payload = {"source_path": "x.pdf", "source_md5": "m5", "schema_version": 1,
               "context": {"authority_for": ["client_name"]}, "tier_a": [], "not_present": []}

    r = _run_write_doc_json(["--out", str(out_dir)], json.dumps(payload))

    assert r.returncode == 0, r.stderr
    assert "WARNING" in r.stderr
    assert "client_name" in r.stderr
    assert r.stdout.strip().startswith("OK ")
    assert (out_dir / (safe_doc_name("x.pdf") + ".json")).exists()


def test_write_doc_json_invalid_json_exits_2(tmp_path):
    out_dir = tmp_path / "out"

    r = _run_write_doc_json(["--out", str(out_dir)], "{not valid json")

    assert r.returncode == 2
    assert "ERROR" in r.stderr
    assert not out_dir.exists()


# ---------------------------------------------------------------------------
# scripts/render_clip.py (subprocess)
# ---------------------------------------------------------------------------


def test_render_clip_whole_page(tmp_path):
    if not SAMPLE_PDF.exists():
        pytest.skip("PDF de mostra no disponible")
    out_png = tmp_path / "whole.png"

    r = subprocess.run(
        [sys.executable, str(RENDER_CLIP), str(SAMPLE_PDF), "--page", "1", "--out", str(out_png)],
        capture_output=True, text=True,
    )

    assert r.returncode == 0, r.stderr
    assert r.stdout.strip().startswith("OK PNG ")
    assert out_png.exists()


def test_render_clip_with_fraction_clip(tmp_path):
    if not SAMPLE_PDF.exists():
        pytest.skip("PDF de mostra no disponible")
    out_png = tmp_path / "clip.png"

    r = subprocess.run(
        [sys.executable, str(RENDER_CLIP), str(SAMPLE_PDF), "--page", "1",
         "--clip", "0.0,0.30,0.5,0.62", "--dpi", "150", "--out", str(out_png)],
        capture_output=True, text=True,
    )

    assert r.returncode == 0, r.stderr
    assert out_png.exists()


def test_render_clip_grid_generates_one_file_per_cell(tmp_path):
    if not SAMPLE_PDF.exists():
        pytest.skip("PDF de mostra no disponible")
    out_png = tmp_path / "grid.png"

    r = subprocess.run(
        [sys.executable, str(RENDER_CLIP), str(SAMPLE_PDF), "--page", "1", "--grid", "2x2", "--out", str(out_png)],
        capture_output=True, text=True,
    )

    assert r.returncode == 0, r.stderr
    lines = [l for l in r.stdout.splitlines() if l.strip()]
    assert len(lines) == 4
    for row in range(2):
        for col in range(2):
            assert (tmp_path / f"grid-r{row}c{col}.png").exists()


def test_render_clip_page_out_of_range_exits_2(tmp_path):
    if not SAMPLE_PDF.exists():
        pytest.skip("PDF de mostra no disponible")
    out_png = tmp_path / "bad.png"

    r = subprocess.run(
        [sys.executable, str(RENDER_CLIP), str(SAMPLE_PDF), "--page", "999", "--out", str(out_png)],
        capture_output=True, text=True,
    )

    assert r.returncode == 2
    assert "ERROR" in r.stderr
    assert not out_png.exists()
