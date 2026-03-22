"""Word document miner (.doc, .docx) -- extracts signals from Word files."""

from __future__ import annotations

import logging
from pathlib import Path

from ..base import BaseMiner
from ..models import Signal, SignalType
from ._detection import run_all_detectors

logger = logging.getLogger(__name__)


class DocxMiner(BaseMiner):
    """Extract signals from Word documents."""

    def can_mine(self, file_path: Path) -> bool:
        return file_path.suffix.lower() in ('.doc', '.docx')

    def mine(self, file_path: Path) -> list[Signal]:
        signals: list[Signal] = []

        text = self._extract_docx_text(file_path)
        if text:
            rel_path = self._relative_path(file_path)
            signals.extend(run_all_detectors(text, rel_path, confidence_offset=0.0))

        try:
            signals.extend(self._extract_images_docx(file_path))
        except Exception:
            logger.debug("Image extraction failed for %s", file_path.name)

        return signals

    def _extract_images_docx(self, file_path: Path) -> list[Signal]:
        """Extract embedded images from Word document."""
        from docx import Document

        signals: list[Signal] = []
        rel = self._relative_path(file_path)
        out_dir = file_path.parent / "validation" / "mined_images"

        doc = Document(str(file_path))
        img_idx = 0
        max_images = 10
        for rel_item in doc.part.rels.values():
            if img_idx >= max_images:
                break
            if "image" not in rel_item.reltype:
                continue
            try:
                blob = rel_item.target_part.blob
                if not blob or len(blob) < 5000:  # Skip tiny images (< 5KB)
                    continue
                content_type = rel_item.target_part.content_type or ""
                ext = "png"
                if "jpeg" in content_type or "jpg" in content_type:
                    ext = "jpg"
                elif "gif" in content_type:
                    ext = "gif"
                elif "emf" in content_type or "wmf" in content_type:
                    continue  # Skip Windows metafiles
                out_dir.mkdir(parents=True, exist_ok=True)
                stem = file_path.stem
                dest = out_dir / f"{stem}_img{img_idx}.{ext}"
                dest.write_bytes(blob)
                signals.append(Signal(
                    type=SignalType.IMAGE,
                    label=f"embedded_image_{img_idx}",
                    value=str(dest),
                    source_file=rel,
                    source_location=f"document image {img_idx}",
                    extraction_method="embedded_image",
                    confidence=0.30,
                ))
                img_idx += 1
            except Exception:
                logger.debug("Could not extract image %d from %s", img_idx, file_path.name)

        return signals

    @staticmethod
    def _extract_doc_legacy(file_path: Path) -> str | None:
        """Extract text from legacy .doc files using antiword or libreoffice."""
        import subprocess

        # Try antiword first (fast, lightweight)
        try:
            result = subprocess.run(
                ["antiword", str(file_path)],
                capture_output=True, text=True, timeout=10,
            )
            if result.returncode == 0 and result.stdout.strip():
                logger.debug("Extracted .doc via antiword: %s (%d chars)", file_path.name, len(result.stdout))
                return result.stdout
        except FileNotFoundError:
            pass  # antiword not installed

        # Fallback: libreoffice conversion to txt
        try:
            import tempfile
            with tempfile.TemporaryDirectory() as tmpdir:
                subprocess.run(
                    ["libreoffice", "--headless", "--convert-to", "txt:Text", "--outdir", tmpdir, str(file_path)],
                    capture_output=True, timeout=30,
                )
                txt_path = Path(tmpdir) / f"{file_path.stem}.txt"
                if txt_path.exists():
                    text = txt_path.read_text(encoding='utf-8', errors='replace')
                    logger.debug("Extracted .doc via libreoffice: %s (%d chars)", file_path.name, len(text))
                    return text
        except FileNotFoundError:
            pass  # libreoffice not installed

        logger.warning("Cannot extract .doc text (no antiword or libreoffice): %s", file_path.name)
        return None

    @staticmethod
    def _extract_docx_text(file_path: Path) -> str | None:
        """Extract text from paragraphs and tables in a Word document.

        For .doc (legacy), falls back to antiword/libreoffice.
        """
        try:
            from docx import Document
        except ImportError:
            logger.debug("python-docx not installed, skipping Word extraction")
            return None

        try:
            doc = Document(str(file_path))
        except Exception as exc:
            if file_path.suffix.lower() == '.doc':
                return DocxMiner._extract_doc_legacy(file_path)
            logger.debug("Could not open Word doc %s: %s", file_path, exc)
            return None

        parts: list[str] = []

        # Extract paragraph text
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                parts.append(text)

        # Extract table cell text
        for table in doc.tables:
            for row in table.rows:
                row_cells: list[str] = []
                for cell in row.cells:
                    cell_text = cell.text.strip()
                    if cell_text:
                        row_cells.append(cell_text)
                if row_cells:
                    # Join cells with tab so label-value detection can find pairs
                    parts.append('\t'.join(row_cells))

        return '\n'.join(parts) if parts else None
