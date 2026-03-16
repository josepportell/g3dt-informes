#!/usr/bin/env python3
"""
Template paragraph migrator for G3DT report improvements.

Modifies the docx template to improve paragraph matching quality
from ~93% to ~96% by updating template text to better match
the reference document patterns.

Usage:
    cd /home/josep/projects/claudecode-job/clients/g3dt
    .venv/bin/python -m automation.template_migrator

Author: Eficients.cat
Date: 2026-02-08
"""

from copy import deepcopy
from pathlib import Path

from docx import Document

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------
BASE_DIR = Path(__file__).resolve().parent.parent
TEMPLATE_PATH = BASE_DIR / "templates" / "g3dt-jinja-template.docx"
OUTPUT_PATH = TEMPLATE_PATH  # Overwrite in-place (backup first)
BACKUP_PATH = TEMPLATE_PATH.with_suffix('.docx.backup-pre-paragraphs')


# ---------------------------------------------------------------------------
# Migration definitions
# ---------------------------------------------------------------------------

PARAGRAPH_MIGRATIONS = [
    {
        'id': 'P54-sol·licitant',
        'paragraph_index': 54,
        'old_contains': 'sol·licitant',
        'new_text': (
            "Segons ens indica el sol·licitant, el SR. {{ architect_name_upper }}, "
            "de l'{{ architect_company }}, en nom de {{ client }}, es vol valorar "
            "les característiques geològiques i geotècniques d'una zona on es preveu "
            "la construcció d'un {{ building_type_lower }}."
        ),
    },
    {
        'id': 'P60-ubicacio',
        'paragraph_index': 60,
        'old_contains': "es situarà",
        'new_text': (
            "L'edificació que es preveu construir es situarà "
            "{{ location_sentence }}."
        ),
    },
    {
        'id': 'P101-adjacent-north',
        'paragraph_index': 101,
        'old_contains': 'nord',
        'new_text': '{{ adjacent_north_fmt }}',
    },
    {
        'id': 'P102-adjacent-south',
        'paragraph_index': 102,
        'old_contains': 'sud',
        'new_text': '{{ adjacent_south_fmt }}',
    },
    {
        'id': 'P103-adjacent-east',
        'paragraph_index': 103,
        'old_contains': 'est',
        'new_text': '{{ adjacent_east_fmt }}',
    },
    {
        'id': 'P104-adjacent-west',
        'paragraph_index': 104,
        'old_contains': 'oest',
        'new_text': '{{ adjacent_west_fmt }}',
    },
    {
        'id': 'P108-acces',
        'paragraph_index': 108,
        'old_contains': "treballs de camp",
        'new_text': (
            "El dia dels treballs de camp es realitza l'entrada a la zona "
            "d'estudi a través del {{ access_street }}."
        ),
    },
    {
        'id': 'P279-erosio',
        'paragraph_index': 279,
        'old_contains': "no s'han detectat marques",
        'new_text': (
            "Degut a que es tracta d'un solar {{ site_condition }}, no s'han "
            "detectat marques i/o indicis de processos d'erosió relacionats amb "
            "l'escolament hídric superficial, ni es preveu que apareguin."
        ),
    },
    {
        'id': 'P281-curs-aigua',
        'paragraph_index': 281,
        'old_contains': "curs d'aigua",
        'new_text': (
            "Tampoc es detecta cap curs d'aigua superficial que pugui "
            "afectar a la zona en estudi."
        ),
    },
    {
        'id': 'P283-hidrogeologia-title',
        'paragraph_index': 283,
        'old_contains': 'Hidrogeologia subterr',
        'new_text': '3.3.2. Hidrogeologia subterrània i geotèrmia',
    },
    {
        'id': 'P436-erosio-conclusions',
        'paragraph_index': 436,
        'old_contains': "no s'han detectat marques",
        'new_text': (
            "Degut a que es tracta d'un solar {{ site_condition }}, no s'han "
            "detectat marques i/o indicis de processos d'erosió relacionats amb "
            "l'escolament hídric superficial, ni es preveu que apareguin."
        ),
    },
    {
        'id': 'P438-curs-aigua-conclusions',
        'paragraph_index': 438,
        'old_contains': "curs d'aigua",
        'new_text': (
            "Tampoc es detecta cap curs d'aigua superficial que pugui "
            "afectar a la zona en estudi."
        ),
    },
]

TABLE_MIGRATIONS = [
    {
        'id': 'T0-R1-C0-superficie',
        'table_index': 0,
        'row_index': 1,
        'cell_index': 0,
        'old_contains': 'Superfície de la parcel·la',
        'new_text': 'Superfície de la parcel·la segons plànols cadastrals  (m2)',
    },
]


# ---------------------------------------------------------------------------
# Migration engine
# ---------------------------------------------------------------------------

def _replace_paragraph_text(paragraph, new_text: str) -> bool:
    """
    Replace all text in a paragraph while preserving the formatting
    of the first run.

    Returns True if replacement was made.
    """
    if not paragraph.runs:
        # No runs — add text directly
        paragraph.text = new_text
        return True

    # Save formatting from first run
    first_run = paragraph.runs[0]
    saved_rpr = deepcopy(first_run._r.get_or_add_rPr())

    # Clear all existing runs
    for run in paragraph.runs:
        run._r.getparent().remove(run._r)

    # Add new run with saved formatting
    from docx.oxml.ns import qn as _qn
    from docx.oxml import OxmlElement
    new_r = OxmlElement('w:r')
    new_r.append(saved_rpr)
    new_t = OxmlElement('w:t')
    new_t.set(_qn('xml:space'), 'preserve')
    new_t.text = new_text
    new_r.append(new_t)
    paragraph._p.append(new_r)

    return True


def _replace_cell_text(cell, new_text: str) -> bool:
    """
    Replace text in a table cell while preserving formatting.

    Returns True if replacement was made.
    """
    if cell.paragraphs:
        return _replace_paragraph_text(cell.paragraphs[0], new_text)
    cell.text = new_text
    return True


def _normalize_apostrophes(text: str) -> str:
    """Normalize curly apostrophes to straight for comparison."""
    return text.replace('\u2019', "'").replace('\u2018', "'")


def _find_paragraph_near_index(paragraphs, target_index: int, old_contains: str, search_range: int = 10):
    """
    Find a paragraph near the target index that contains the expected text.

    Searches target_index first, then expands search ±search_range.
    Normalizes curly apostrophes for comparison.
    Returns (actual_index, paragraph) or (None, None).
    """
    needle = _normalize_apostrophes(old_contains.lower())

    # Try exact index first
    if 0 <= target_index < len(paragraphs):
        text = _normalize_apostrophes(paragraphs[target_index].text.lower())
        if needle in text:
            return target_index, paragraphs[target_index]

    # Search nearby
    for offset in range(1, search_range + 1):
        for idx in [target_index + offset, target_index - offset]:
            if 0 <= idx < len(paragraphs):
                text = _normalize_apostrophes(paragraphs[idx].text.lower())
                if needle in text:
                    return idx, paragraphs[idx]

    return None, None


def migrate_template(
    input_path: Path | None = None,
    output_path: Path | None = None,
    dry_run: bool = False,
) -> dict:
    """
    Apply all migrations to the template.

    Args:
        input_path: Source template (default: TEMPLATE_PATH)
        output_path: Where to save (default: OUTPUT_PATH, overwrites)
        dry_run: If True, don't save — just report what would change

    Returns:
        dict with 'applied', 'skipped', 'errors' lists
    """
    input_path = input_path or TEMPLATE_PATH
    output_path = output_path or OUTPUT_PATH

    if not input_path.exists():
        return {'applied': [], 'skipped': [], 'errors': [f'Template not found: {input_path}']}

    doc = Document(str(input_path))
    paragraphs = doc.paragraphs
    tables = doc.tables

    result = {'applied': [], 'skipped': [], 'errors': []}

    # Apply paragraph migrations
    for mig in PARAGRAPH_MIGRATIONS:
        mid = mig['id']
        target_idx = mig['paragraph_index']
        old_contains = mig['old_contains']
        new_text = mig['new_text']

        actual_idx, para = _find_paragraph_near_index(paragraphs, target_idx, old_contains)

        if para is None:
            result['skipped'].append(f'{mid}: paragraph not found (expected ~P{target_idx}, looking for "{old_contains}")')
            continue

        old_text = para.text
        if dry_run:
            result['applied'].append(f'{mid}: P{actual_idx} would change from "{old_text[:80]}..." to "{new_text[:80]}..."')
        else:
            _replace_paragraph_text(para, new_text)
            result['applied'].append(f'{mid}: P{actual_idx} updated')

    # Apply table migrations
    for mig in TABLE_MIGRATIONS:
        mid = mig['id']
        t_idx = mig['table_index']
        r_idx = mig['row_index']
        c_idx = mig['cell_index']
        old_contains = mig['old_contains']
        new_text = mig['new_text']

        if t_idx >= len(tables):
            result['skipped'].append(f'{mid}: table {t_idx} not found (only {len(tables)} tables)')
            continue

        table = tables[t_idx]
        if r_idx >= len(table.rows):
            result['skipped'].append(f'{mid}: row {r_idx} not found in table {t_idx}')
            continue

        row = table.rows[r_idx]
        if c_idx >= len(row.cells):
            result['skipped'].append(f'{mid}: cell {c_idx} not found in row {r_idx} of table {t_idx}')
            continue

        cell = row.cells[c_idx]
        cell_text = cell.text

        if _normalize_apostrophes(old_contains.lower()) not in _normalize_apostrophes(cell_text.lower()):
            result['skipped'].append(f'{mid}: expected "{old_contains}" in cell, got "{cell_text[:60]}"')
            continue

        if dry_run:
            result['applied'].append(f'{mid}: T{t_idx}R{r_idx}C{c_idx} would change from "{cell_text[:60]}" to "{new_text[:60]}"')
        else:
            _replace_cell_text(cell, new_text)
            result['applied'].append(f'{mid}: T{t_idx}R{r_idx}C{c_idx} updated')

    # Save
    if not dry_run and result['applied']:
        # Backup first
        if not BACKUP_PATH.exists():
            import shutil
            shutil.copy2(str(input_path), str(BACKUP_PATH))
            print(f'Backup saved to: {BACKUP_PATH}')

        doc.save(str(output_path))
        print(f'Template saved to: {output_path}')

    return result


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main():
    """Run template migration."""
    import argparse

    parser = argparse.ArgumentParser(description='Migrate G3DT template paragraphs')
    parser.add_argument('--dry-run', action='store_true', help='Show what would change without saving')
    parser.add_argument('--input', type=Path, default=None, help='Input template path')
    parser.add_argument('--output', type=Path, default=None, help='Output template path')
    args = parser.parse_args()

    print(f'Template: {args.input or TEMPLATE_PATH}')
    print(f'Mode: {"DRY RUN" if args.dry_run else "LIVE"}')
    print()

    result = migrate_template(
        input_path=args.input,
        output_path=args.output,
        dry_run=args.dry_run,
    )

    if result['applied']:
        print(f'Applied ({len(result["applied"])}):')
        for item in result['applied']:
            print(f'  ✓ {item}')

    if result['skipped']:
        print(f'\nSkipped ({len(result["skipped"])}):')
        for item in result['skipped']:
            print(f'  ⚠ {item}')

    if result['errors']:
        print(f'\nErrors ({len(result["errors"])}):')
        for item in result['errors']:
            print(f'  ✗ {item}')

    total = len(result['applied']) + len(result['skipped'])
    print(f'\nSummary: {len(result["applied"])}/{total} migrations applied')


if __name__ == '__main__':
    main()
