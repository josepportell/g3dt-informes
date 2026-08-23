"""Fase 4a (via A) — Lectors deterministes de les 5 plantilles G3.

Llegeix les cel·les/posicions EXACTES on mira l'Eva (calibrades amb la lectura d'or
dels 8 projectes, `docs/golden-read/`; vegeu el Pas 2 del skill g3dt-llegir-projecte):

  1. Pressupost PDF (m4PRO ERP): p.1 blocs CLIENT/OBRA ordenats per (y,x); p.2 CTE + previstos
  2. Fitxa de camp (DADES PER ANAR A CAMP*.xlsx): fitxa!C6/C7/C8/F8/C13/F13/C16/C38/F38
  3. Comanda de laboratori (.xls): Hoja1!N18-N23 (bloc OBRA), AH23, fila 35+ (mostres);
     el bloc N11-N14 (SOL.LICITANT = G3) s'emet com a NOT_client_name
  4. PLAN_COST (.xlsx): 'Plan Cost'!E9/E21 + OFERTA!B21/B6/B14 (G3/G27 subcontractes)
  5. Excel DPSH (.xls): fulls P-*, executats = fulls amb dades a la columna C; peu B80

Cada senyal porta cel·la/posició i cita literal. Els CLIENT: de pressupost i fitxa són el
SOL·LICITANT (regla G3): s'emeten amb confiança baixa i nota — mai autoritat de client_name.
Cost 0 (sense LLM), re-executable sempre (ANALISI §7.3, etapa 4a).

Ús:
    .venv/bin/python -m automation.g3_templates "/mnt/c/claude/g3dt/projectes/4001612 BELL-LLOC" [--out FILE]

    from automation.g3_templates import read_project
    resultat = read_project(Path(...))   # {'documents': [...], 'concepts': {...}}
"""

from __future__ import annotations

import hashlib
import json
import re
import sys
import warnings
from datetime import datetime
from pathlib import Path

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

EXCLUDE_DIRS = {
    "validation", "_validation", "mined_images", "concept_probes", "msg_attachments",
    "PDF-V0", "PDF V0", "PDF_V0", "LLETRA",
}
EXCLUDE_NAME_PARTS = ("_generated", "informe", "PORTADA", "AUDIT", "~$")

# Ordre de preferència de fonts per concepte (regles del skill, Pas 3). Defecte: ordre de lectura.
PRIORITY = {
    "expedient": ["comanda_lab_g3", "dpsh_excel"],
    "field_date": ["fitxa_camp_g3", "comanda_lab_g3"],
    "municipality": ["comanda_lab_g3", "pressupost_g3", "plan_cost_g3", "fitxa_camp_g3"],
    "street_address": ["pressupost_g3", "fitxa_camp_g3", "comanda_lab_g3"],
    "num_dpsh_tests": ["dpsh_excel", "pressupost_g3", "plan_cost_g3", "fitxa_camp_g3"],
    "num_sondeigs": ["pressupost_g3", "plan_cost_g3"],
    "building_type": ["comanda_lab_g3", "plan_cost_g3"],
    "lab_sample_id": ["comanda_lab_g3"],
    "lab_depth": ["comanda_lab_g3"],
    "lab_location": ["comanda_lab_g3"],
}
_DOC_ORDER = ["pressupost_g3", "fitxa_camp_g3", "comanda_lab_g3", "plan_cost_g3", "dpsh_excel"]


def _sig(concept, value, location, quote, confidence=0.9, note=None):
    s = {"concept_id": concept, "value": value, "location": location,
         "quote": str(quote)[:300], "confidence": confidence}
    if note:
        s["note"] = note
    return s


def _clean(v):
    return " ".join(str(v).split()) if v is not None else ""


# ---------------------------------------------------------------- 1. Pressupost PDF

def read_pressupost_pdf(path: Path):
    import fitz

    doc = fitz.open(path)
    creator = doc.metadata.get("creator") or ""
    if "G3 DESENVOLUPAMENT TERRITORIAL" not in creator:
        return None
    signals, warnings = [], []
    p1 = doc[0]
    # Ordena els blocs per (y, x): l'ordre natural del PDF barreja CLIENT i OBRA
    blocks = sorted(p1.get_text("blocks"), key=lambda b: (round(b[1]), round(b[0])))
    lines = []
    for b in blocks:
        lines.extend(_clean(x) for x in b[4].splitlines() if _clean(x))

    def lines_after(label_re, n=4):
        for i, ln in enumerate(lines):
            if re.match(label_re, ln, re.I):
                out = []
                rest = re.sub(label_re, "", ln, flags=re.I).strip()
                if rest:
                    out.append(rest)
                for nxt in lines[i + 1:]:
                    if re.match(r"^(CLIENT[E]?\s*:|OBRA\s*:|Data\b|Fecha\b)", nxt, re.I):
                        break
                    out.append(nxt)
                    if len(out) >= n:
                        break
                return out
        return []

    client = lines_after(r"^CLIENT[E]?\s*:\s*")
    if client:
        name = client[0]
        note = "= SOL·LICITANT del pressupost (sovint l'arquitecte); NO és autoritat per a client_name (regla G3)"
        signals.append(_sig("client_name", name, "p.1 bloc CLIENT (y,x)", " / ".join(client[:3]), 0.3, note))
        if re.search(r"ARQUITECT", name, re.I):
            signals.append(_sig("architect_name", name, "p.1 bloc CLIENT", name, 0.5,
                                "despatx del sol·licitant; persona si es pot (regla persona/despatx)"))
    obra = lines_after(r"^OBRA\s*:\s*")
    if obra:
        # El bloc OBRA té 3 formes reals (lectura d'or): [encàrrec, adreça, municipi] (Bell-lloc,
        # Vilanova), [adreça+CP+municipi en 1 línia, encàrrec] (Linyola), [encàrrec, municipi] (Rubí).
        def is_addr(ln):
            return bool(re.search(r"\d", ln)) and bool(
                re.search(r"(?:^C/|\bC/|\bAV|\bPL\b|\bCTRA|\bURB|\bCARRER|\bCALLE|\bPASSEIG|"
                          r"\bCAM[IÍ]|\bPOL|,\s*\d|Nº|\b\d{5}\b)", ln, re.I))

        def is_encarrec(ln):
            return bool(re.search(r"ESTUDI|ESTUDIO|PRESSUPOST|PRESUPUESTO|GEOTEC", ln, re.I))

        quote = " / ".join(obra[:4])
        street = next((ln for ln in obra if is_addr(ln)), None)
        municipality = None
        if street:
            m = re.match(r"(.+?)\s*[-–]?\s*\b(\d{5})\b\s*[-–]?\s*(.+)$", street)
            if m:  # adreça + CP + municipi a la mateixa línia (Linyola)
                municipality = m.group(3).strip()
                signals.append(_sig("street_address", m.group(1).strip(",- "),
                                    "p.1 bloc OBRA (línia amb CP)", quote))
            else:
                signals.append(_sig("street_address", street, "p.1 bloc OBRA", quote))
                after = obra[obra.index(street) + 1:]
                municipality = next((ln for ln in after if not is_addr(ln) and not is_encarrec(ln)), None)
        if municipality is None:
            municipality = next((ln for ln in obra if not is_addr(ln) and not is_encarrec(ln)), None)
        if municipality:
            signals.append(_sig("municipality", municipality, "p.1 bloc OBRA", quote, 0.85,
                                "forma curta G3; la forma oficial llarga mana (regla municipi)"))
        if not street:
            warnings.append("bloc OBRA sense línia d'adreça (forma Rubí): street_address no emès")
    if not client or not obra:
        warnings.append("bloc CLIENT o OBRA no localitzat a p.1")

    full = "\n".join(doc[i].get_text() for i in range(min(4, doc.page_count)))
    m = re.search(r"\b(\d{2}·\d{4})\b", full)
    if m:
        signals.append(_sig("expedient_comercial", m.group(1), "p.1 peu / capçaleres", m.group(0), 0.9,
                            "codi comercial m4PRO, NO és l'expedient"))
    m = re.search(r"Tip(?:us|o)\s+d[e’'l\s]*edifici[o]?\s*[.:]?\s*(C[-\s]?\d)", full, re.I)
    if m:
        signals.append(_sig("cte_edificacio", m.group(1).replace(" ", ""), "p.2", _clean(m.group(0))))
    m = re.search(r"Tip(?:us|o)\s+de\s+[Tt]erren[yo]\s*[.:]?\s*(T[-\s]?\d)", full, re.I)
    if m:
        signals.append(_sig("cte_sol", m.group(1).replace(" ", ""), "p.2", _clean(m.group(0))))
    m = re.search(r"(\d+)\s*(?:assaigs?|ensayos?)\s+de\s+penetraci[óo]n?\s+din[àá]mica", full, re.I)
    if m:
        signals.append(_sig("num_dpsh_tests", int(m.group(1)), "p.2 campanya", _clean(m.group(0)), 0.7,
                            "PREVISTOS; si l'Excel DPSH en diu un altre, mana l'Excel"))
    m = re.search(r"(\d+)\s*[Ss]onde(?:ig|o)s?\s+a\s+rotaci[óo]n?", full, re.I)
    if m:
        signals.append(_sig("num_sondeigs", int(m.group(1)), "p.2 campanya", _clean(m.group(0)), 0.7, "previstos"))

    return {"document_type": "pressupost_g3", "mod_date": doc.metadata.get("modDate") or "",
            "tier_a": signals, "warnings": warnings}


# ---------------------------------------------------------------- 2. Fitxa de camp

def read_fitxa_camp(path: Path):
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    name = next((n for n in wb.sheetnames if n.strip().lower() == "fitxa"), None)
    if name is None:
        return None
    ws = wb[name]
    if _clean(ws["B2"].value).upper() != "DADES PER ANAR A CAMP":
        return None
    signals, warnings = [], []

    def cell(ref, concept, conf=0.9, note=None, label=""):
        v = ws[ref].value
        if v is None or _clean(v) == "":
            return None
        signals.append(_sig(concept, _clean(v), f"fitxa!{ref}" + (f" ({label})" if label else ""),
                            _clean(v), conf, note))
        return v

    cell("C6", "client_name", 0.3, "= sol·licitant (regla G3); NO autoritat de client_name", "CLIENT")
    cell("C7", "street_address", 0.8, "C7 porta adreça + municipi a la mateixa cel·la", "ADREÇA OBRA")
    cell("C8", "contact_person", 0.8, label="PERSONA DE CONTACTE")
    cell("F8", "contact_phone", 0.8)
    cell("C13", "previsio_camp", 0.7, label="PREVISIÓ DE TREBALL DE CAMP")
    cell("F13", "previsio_camp_unitats", 0.7, "p. ex. '2 (P)', '5P+2S' — previstos, no executats")
    cell("C16", "edificacio_existent", 0.6, label="EDIFICACIÓ EXISTENT ?")
    cell("C38", "lab_field_company", 0.8, label="EMPRESA QUE FA LA FEINA")
    f38 = ws["F38"].value
    if isinstance(f38, datetime):
        signals.append(_sig("field_date", f38.date().isoformat(), "fitxa!F38 (E38 'dies de camp')",
                            repr(f38), 0.95, "primer dia de camp (autoritat A)"))
    elif f38 is not None and _clean(f38):
        signals.append(_sig("field_date", _clean(f38), "fitxa!F38", _clean(f38), 0.7, "no és tipus data"))
    else:
        warnings.append("F38 (data de camp) buit: usar annex DPSH / comanda / albarà / fotos")

    return {"document_type": "fitxa_camp_g3", "tier_a": signals, "warnings": warnings}


# ---------------------------------------------------------------- 3. Comanda de laboratori

def read_comanda(path: Path):
    import xlrd

    book = xlrd.open_workbook(str(path))
    try:
        sh = book.sheet_by_name("Hoja1")
    except xlrd.XLRDError:
        sh = book.sheet_by_index(0)
    c3 = _clean(sh.cell_value(2, 2)).upper() if sh.nrows > 2 and sh.ncols > 2 else ""
    if "PETICI" not in c3 or "LABORATORI" not in c3:
        return None
    signals, warnings = [], []
    N = 13  # columna N

    def val(r0, c0):
        return sh.cell_value(r0, c0) if r0 < sh.nrows and c0 < sh.ncols else None

    def dateval(r0, c0):
        v = val(r0, c0)
        if isinstance(v, float) and v > 1000:
            try:
                return xlrd.xldate_as_datetime(v, book.datemode).date().isoformat()
            except Exception:
                return None
        return _clean(v) or None

    exp = val(18, N)  # N19
    if isinstance(exp, float):
        exp = str(int(exp))
    if exp and _clean(exp):
        signals.append(_sig("expedient", _clean(exp), "Hoja1!N19 (G19 NÚM. D'EXPEDIENT)", repr(val(18, N)), 0.95))
    if _clean(val(17, N)):
        signals.append(_sig("building_type", _clean(val(17, N)), "Hoja1!N18 (G18 OBRA)", _clean(val(17, N)), 0.7))
    if _clean(val(19, N)):
        signals.append(_sig("street_address", _clean(val(19, N)), "Hoja1!N20 (G20 ADREÇA, bloc OBRA)", _clean(val(19, N))))
    if _clean(val(20, N)):
        signals.append(_sig("municipality", _clean(val(20, N)), "Hoja1!N21 (G21 POBLACIÓ, bloc OBRA)", _clean(val(20, N))))
    presa = dateval(22, N)  # N23
    if presa:
        signals.append(_sig("field_date", presa, "Hoja1!N23 (DATA DE PRESA)", str(val(22, N)), 0.7,
                            "dia de la mostra/sondeig, NO necessàriament el primer dia de camp"))
    sol = dateval(22, 33)  # AH23
    if sol:
        signals.append(_sig("data_sollicitud_lab", sol, "Hoja1!AH23", str(val(22, 33)), 0.6))
    # mostres: fila 35+ (idx 34); C=id, J/L=cotes, AD=nota
    for r in range(34, min(45, sh.nrows)):
        mid = _clean(val(r, 2))
        if not mid:
            continue
        signals.append(_sig("lab_sample_id", mid, f"Hoja1!C{r + 1}", mid, 0.7,
                            "etiqueta de la comanda; l'annex de l'Eva mana per a l'etiqueta de l'informe"))
        ini, fin = val(r, 9), val(r, 11)
        if ini is not None and fin is not None and str(ini) != "" and str(fin) != "":
            signals.append(_sig("lab_depth", f"{ini} - {fin}", f"Hoja1!J{r + 1}/L{r + 1}",
                                f"INICIAL {ini} / FINAL {fin}", 0.7, "el GTL mana si discrepa"))
        m = re.search(r"\(([^)]+)\)", mid)
        if m:
            signals.append(_sig("lab_location", m.group(1), f"Hoja1!C{r + 1} (parèntesi)", mid, 0.7))
        if _clean(val(r, 29)):
            signals.append(_sig("lab_tests_note", _clean(val(r, 29)), f"Hoja1!AD{r + 1}", _clean(val(r, 29)), 0.6))
    # bloc SOL.LICITANT (files 11-15) = G3: NOT_client explícit
    g3name = _clean(val(10, N))
    if g3name:
        signals.append(_sig("NOT_client_name", g3name, "Hoja1!N11-N14 (bloc DADES DEL SOL.LICITANT)",
                            f"{g3name} | {_clean(val(11, N))}", 1.0, "G3 = sol·licitant del laboratori, mai client"))
    return {"document_type": "comanda_lab_g3", "tier_a": signals, "warnings": warnings}


# ---------------------------------------------------------------- 4. PLAN_COST

def read_plan_cost(path: Path):
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    name = next((n for n in wb.sheetnames if n.strip().lower() == "plan cost"), None)
    if name is None:
        return None
    ws = wb[name]
    if not _clean(ws["B2"].value).upper().startswith("PLAN COST"):
        return None
    signals, warnings = [], []
    e9 = _clean(ws["E9"].value)
    if e9:
        signals.append(_sig("building_type", e9, "'Plan Cost'!E9 (B9 DESCRIPCIÓ TITÒL)", e9, 0.6,
                            "format 'EG HAB UNIF {MUNICIPI}': expandir tipus; el final és el municipi"))
        m = re.match(r"EG\s+(.*?)\s+([A-ZÀ-Ü' .-]+)$", e9)
        if m:
            signals.append(_sig("municipality", m.group(2).strip(), "'Plan Cost'!E9 (final)", e9, 0.6))
    if _clean(ws["E21"].value):
        signals.append(_sig("tecnic", _clean(ws["E21"].value), "'Plan Cost'!E21", _clean(ws["E21"].value), 0.6))
    of = next((n for n in wb.sheetnames if n.strip().upper() == "OFERTA"), None)
    if of:
        o = wb[of]
        for ref, concept, note in (("B21", "num_dpsh_tests", "unitats DPSH ofertades (previstes)"),
                                   ("B6", "sondeig_ml", "ml de perforació"),
                                   ("B14", "num_spt", "assaigs SPT previstos")):
            v = o[ref].value
            if isinstance(v, (int, float)) and v:
                signals.append(_sig(concept, int(v) if float(v).is_integer() else v,
                                    f"OFERTA!{ref}", str(v), 0.6, note))
        for ref in ("G3", "G27"):
            if _clean(o[ref].value):
                signals.append(_sig("subcontractes_hint", _clean(o[ref].value), f"OFERTA!{ref}",
                                    _clean(o[ref].value), 0.4, "plantilla: mai senyal fort"))
    else:
        warnings.append("full OFERTA absent")
    return {"document_type": "plan_cost_g3", "tier_a": signals, "warnings": warnings}


# ---------------------------------------------------------------- 5. Excel DPSH

def read_dpsh_excel(path: Path):
    import xlrd

    book = xlrd.open_workbook(str(path))
    psheets = [n for n in book.sheet_names() if re.match(r"^P[-.]?\d", n.strip())]
    if not psheets:
        return None
    signals, executed = [], []
    for n in psheets:
        sh = book.sheet_by_name(n)
        has_data = any(
            isinstance(sh.cell_value(r, 2), (int, float)) and str(sh.cell_value(r, 2)) != ""
            for r in range(16, min(90, sh.nrows))
        )
        if has_data:
            executed.append(n)
            b80 = _clean(sh.cell_value(79, 1)) if sh.nrows > 79 else ""
            if "ebuig" in b80 or "echa" in b80:
                signals.append(_sig("dpsh_refusal", b80, f"{n}!B80", b80, 0.8))
    signals.insert(0, _sig("num_dpsh_tests", len(executed), "fulls amb dades a la columna C",
                           f"sheets: {executed}", 0.95, "EXECUTATS (autoritat sobre els previstos)"))
    m = re.match(r"^(\d{7})", path.name)
    if m:
        signals.append(_sig("expedient", m.group(1), "nom del fitxer", path.name, 0.7))
    return {"document_type": "dpsh_excel", "tier_a": signals, "warnings": []}


# ---------------------------------------------------------------- escaneig + agregació

_READERS = {".pdf": [read_pressupost_pdf],
            ".xlsx": [read_fitxa_camp, read_plan_cost],
            ".xls": [read_comanda, read_dpsh_excel]}


def read_project(project_dir: Path) -> dict:
    project_dir = Path(project_dir)
    documents, seen_md5 = [], {}
    for p in sorted(project_dir.rglob("*")):
        if not p.is_file() or p.suffix.lower() not in _READERS:
            continue
        rel = p.relative_to(project_dir)
        if any(part in EXCLUDE_DIRS for part in rel.parts):
            continue
        if any(t.lower() in p.name.lower() for t in EXCLUDE_NAME_PARTS):
            continue
        md5 = hashlib.md5(p.read_bytes()).hexdigest()
        if md5 in seen_md5:
            documents.append({"source_path": str(rel), "document_type": "duplicat",
                              "duplicate_of": seen_md5[md5], "tier_a": [], "warnings": []})
            continue
        for reader in _READERS[p.suffix.lower()]:
            try:
                doc = reader(p)
            except Exception as exc:  # un fitxer corrupte no ha d'aturar la resta
                doc = None
                documents.append({"source_path": str(rel), "document_type": "error",
                                  "tier_a": [], "warnings": [f"{reader.__name__}: {exc}"]})
                break
            if doc is not None:
                doc["source_path"] = str(rel)
                documents.append(doc)
                seen_md5[md5] = str(rel)
                break

    concepts: dict[str, list] = {}
    for doc in documents:
        for s in doc.get("tier_a", []):
            cand = dict(s)
            cand["source"] = doc["source_path"]
            cand["document_type"] = doc["document_type"]
            if doc.get("mod_date"):
                cand["mod_date"] = doc["mod_date"]
            concepts.setdefault(s["concept_id"], []).append(cand)
    for cid, cands in concepts.items():
        order = PRIORITY.get(cid, _DOC_ORDER)
        # multi-pas estable: confiança ↓, després modDate ↓ (el pressupost MODF més nou mana),
        # i finalment la prioritat de font per concepte
        cands.sort(key=lambda c: -c.get("confidence", 0))
        cands.sort(key=lambda c: c.get("mod_date", ""), reverse=True)
        cands.sort(key=lambda c: (order.index(c["document_type"]) if c["document_type"] in order
                                  else len(order)))
    return {"project": project_dir.name, "read_on": datetime.now().isoformat(timespec="seconds"),
            "reader": "automation.g3_templates (Fase 4a, determinista)",
            "documents": documents, "concepts": concepts}


def main() -> None:
    args = sys.argv[1:]
    out = None
    if "--out" in args:
        i = args.index("--out")
        out = Path(args[i + 1])
        del args[i:i + 2]
    if len(args) != 1:
        sys.exit("Ús: python -m automation.g3_templates PROJECT_DIR [--out FILE]")
    result = read_project(Path(args[0]))
    ndocs = sum(1 for d in result["documents"] if d["document_type"] not in ("duplicat", "error"))
    payload = json.dumps(result, ensure_ascii=False, indent=1)
    if out:
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(payload, encoding="utf-8")
        print(f"{result['project']}: {ndocs} plantilles G3 llegides, "
              f"{len(result['concepts'])} conceptes → {out}")
    else:
        print(payload)


if __name__ == "__main__":
    main()
