"""
SmartScan Tier 2: Structural fingerprinting.

Classifies files by their actual content, independent of filename.
Uses PyMuPDF for PDFs, openpyxl for .xlsx, and xlrd for .xls files.

This is the heart of SmartScan's robustness — it works with files
named ANYTHING by analyzing what the file IS, not what it's called.
"""

from __future__ import annotations

import logging
from pathlib import Path

from .confidence import tier2_confidence
from .models import ClassificationTier, FileClassification

logger = logging.getLogger(__name__)


# ──────────────────────────────────────────────────────────────
# Role fingerprint definitions
# ──────────────────────────────────────────────────────────────
# Each role has keywords (CA + ES), structural signals, and
# discrimination rules.

ROLE_FINGERPRINTS: dict[str, dict] = {
    'dpsh_field_sheet': {
        'keywords': [
            'N20', 'DPSH', 'penetr', 'profunditat', 'profundidad',
            'cops', 'golpes', 'refús', 'rechazo', 'N.F.',
            'penetròmetre', 'penetrómetro',
        ],
        'page_profile': 'scanned',     # Mostly image (handwritten)
        'max_pages': 10,
        'typical_pages': (1, 5),
    },
    'sondeig_field_sheet': {
        'keywords': [
            'sondeig', 'sondeo', 'rotació', 'rotación',
            'SPT', 'capes', 'capas', 'profunditat', 'profundidad',
            'N.F.', 'roca', 'argila', 'arcilla', 'grava',
        ],
        'page_profile': 'scanned',
        'max_pages': 10,
        'typical_pages': (1, 5),
    },
    'architect_plan': {
        'keywords': [
            'arquitecte', 'arquitecto', 'promotor', 'planta',
            'alçat', 'alzado', 'façana', 'fachada', 'parcel·la',
            'parcela', 'cota', 'escala', 'visat', 'visado',
        ],
        'page_profile': 'vector',       # CAD-exported
        'producer_hints': ['AutoCAD', 'Revit', 'ArchiCAD', 'DWG'],
        'max_pages': 5,
        'typical_pages': (1, 3),
        'orientation': 'landscape',
    },
    'sondeig_annex': {
        'keywords': [
            'sondeig', 'sondeo', 'unitat litològica', 'unidad litológica',
            'profunditat', 'profundidad', 'descripció', 'descripción',
            'SPT', 'columna', 'capa',
        ],
        'page_profile': 'vector',
        'producer_hints': ['FreeHand', 'Illustrator', 'Acrobat'],
        'max_pages': 5,
        'typical_pages': (1, 3),
    },
    'correlation_section': {
        'keywords': [
            'tall', 'corte', 'correlació', 'correlación',
            'secció', 'sección', 'P-1', 'P-2', 'S-1',
        ],
        'page_profile': 'vector',
        'max_pages': 3,
        'typical_pages': (1, 2),
    },
    'situation_plan': {
        'keywords': [
            'situació', 'situación', 'ubicació', 'ubicación',
            'plànol', 'plano', 'emplaçament', 'emplazamiento',
            'escala', 'nord', 'norte',
        ],
        'page_profile': 'vector',
        'max_pages': 3,
        'typical_pages': (1, 2),
    },
    'reference_report': {
        'keywords': [
            'informe', 'geotècnic', 'geotécnico', 'estudi',
            'estudio', 'fonamentació', 'cimentación',
            'promotor', 'objecte', 'objeto',
        ],
        'page_profile': 'text',
        'min_pages': 5,
        'typical_pages': (10, 50),
    },
    'lab_results_pdf': {
        'keywords': [
            'laboratori', 'laboratorio', 'assaig', 'ensayo',
            'sulfat', 'sulfato', 'mg/kg', 'agressivitat',
            'agresividad', 'resultat', 'resultado',
        ],
        'page_profile': 'mixed',
        'max_pages': 10,
        'typical_pages': (1, 5),
    },
    'gtl_report': {
        'keywords': [
            'GTL', 'laboratori', 'laboratorio', 'acreditació',
            'acreditación', 'ENAC', 'assaig', 'ensayo',
        ],
        'page_profile': 'text',
        'max_pages': 15,
        'typical_pages': (2, 10),
    },
    'dpsh_excel': {
        'keywords_excel': [
            'N20', 'profunditat', 'profundidad', 'DPSH',
            'cops', 'golpes', 'refus', 'rechazo',
        ],
        'is_excel': True,
        'depth_pattern': True,  # Values every 0.20m
    },
    'lab_order': {
        'keywords_excel': [
            'comanda', 'pedido', 'laboratori', 'laboratorio',
            'mostra', 'muestra', 'assaig', 'ensayo',
        ],
        'is_excel': True,
    },
}


def classify_tier2(
    project_path: Path,
    entries: list[tuple[str, bool]],
    already_classified: set[str],
) -> list[FileClassification]:
    """
    Classify unclassified files using structural fingerprinting.

    Only processes files NOT already classified by Tier 1.
    Uses PyMuPDF for PDFs, openpyxl/xlrd for Excel, Pillow for images.

    Args:
        project_path: Absolute path to the project folder
        entries: All (relative_path, is_directory) tuples
        already_classified: Set of relative paths already classified

    Returns:
        List of FileClassification for files classified by fingerprinting.
    """
    results: list[FileClassification] = []

    for rel_path, is_dir in entries:
        if rel_path in already_classified or is_dir:
            continue

        abs_path = project_path / rel_path
        ext = abs_path.suffix.lower()

        if ext == '.pdf':
            clf = _fingerprint_pdf(rel_path, abs_path)
            if clf:
                results.append(clf)
        elif ext in ('.xls', '.xlsx'):
            clf = _fingerprint_excel(rel_path, abs_path)
            if clf:
                results.append(clf)
        elif ext in ('.jpg', '.jpeg', '.png'):
            clf = _fingerprint_image(rel_path, abs_path)
            if clf:
                results.append(clf)

    return results


def _fingerprint_pdf(rel_path: str, abs_path: Path) -> FileClassification | None:
    """Analyze a PDF's structure and content to determine its role."""
    try:
        import fitz  # PyMuPDF
    except ImportError:
        logger.warning("PyMuPDF not available for Tier 2 fingerprinting")
        return None

    try:
        doc = fitz.open(str(abs_path))
    except Exception as e:
        logger.warning(f"Cannot open PDF {rel_path}: {e}")
        return None

    try:
        num_pages = len(doc)
        if num_pages == 0:
            return None

        # Extract metadata
        metadata = doc.metadata or {}
        producer = (metadata.get('producer') or '').lower()
        creator = (metadata.get('creator') or '').lower()

        # Analyze first page
        page = doc[0]
        width, height = page.rect.width, page.rect.height
        is_landscape = width > height * 1.1

        # Get text from first pages (up to 3)
        text_sample = ""
        for i in range(min(num_pages, 3)):
            text_sample += doc[i].get_text("text")

        text_lower = text_sample.lower()
        text_len = len(text_sample.strip())

        # Determine page profile (scanned vs vector vs text)
        blocks = page.get_text("blocks")
        img_blocks = sum(1 for b in blocks if b[6] == 1)  # type 1 = image
        text_blocks = sum(1 for b in blocks if b[6] == 0)  # type 0 = text
        total_blocks = img_blocks + text_blocks

        if total_blocks == 0:
            page_profile = 'empty'
        elif img_blocks > 0 and text_blocks == 0:
            page_profile = 'scanned'
        elif img_blocks > text_blocks * 2:
            page_profile = 'scanned'
        elif text_len > 2000:
            page_profile = 'text'
        elif is_landscape:
            page_profile = 'vector'
        else:
            page_profile = 'mixed'

        # Build fingerprint data
        fingerprint: dict = {
            'num_pages': num_pages,
            'page_profile': page_profile,
            'is_landscape': is_landscape,
            'width': round(width, 1),
            'height': round(height, 1),
            'text_length': text_len,
            'producer': producer,
            'creator': creator,
            'img_blocks': img_blocks,
            'text_blocks': text_blocks,
        }

        # Score each role
        best_role: str | None = None
        best_score = 0.0
        alternates: list[dict] = []

        for role_name, fp_def in ROLE_FINGERPRINTS.items():
            if fp_def.get('is_excel'):
                continue

            score = _score_pdf_role(
                role_name, fp_def,
                text_lower, num_pages, page_profile,
                is_landscape, producer, creator,
            )

            if score > 0:
                alternates.append({'role': role_name, 'confidence': round(score, 3)})
                if score > best_score:
                    best_score = score
                    best_role = role_name

        doc.close()

        if best_role and best_score >= 0.3:
            from .confidence import classify_category
            return FileClassification(
                file_path=rel_path,
                role=best_role,
                confidence=best_score,
                tier=ClassificationTier.FINGERPRINT,
                alternate_roles=[a for a in alternates if a['role'] != best_role],
                fingerprint_data=fingerprint,
                category=classify_category(best_score, best_role),
            )

        # Return as unclassified with fingerprint data for debug
        if fingerprint:
            return FileClassification(
                file_path=rel_path,
                role=None,
                confidence=0.0,
                tier=ClassificationTier.FINGERPRINT,
                fingerprint_data=fingerprint,
                category="unknown",
                summary=f"PDF {num_pages}p, {page_profile}, {text_len} chars",
            )

    except Exception as e:
        logger.warning(f"Tier 2 PDF analysis failed for {rel_path}: {e}")
        try:
            doc.close()
        except Exception:
            pass

    return None


def _score_pdf_role(
    role_name: str,
    fp_def: dict,
    text_lower: str,
    num_pages: int,
    page_profile: str,
    is_landscape: bool,
    producer: str,
    creator: str,
) -> float:
    """Score how well a PDF matches a given role's fingerprint."""
    keyword_hits = 0
    keywords = fp_def.get('keywords', [])
    for kw in keywords:
        if kw.lower() in text_lower:
            keyword_hits += 1

    if not keywords or keyword_hits == 0:
        return 0.0

    # Keyword ratio (how many of the expected keywords are present)
    kw_ratio = keyword_hits / len(keywords)
    if kw_ratio < 0.15:
        return 0.0

    # Page profile match
    expected_profile = fp_def.get('page_profile')
    profile_match = (expected_profile == page_profile) if expected_profile else False

    # Page count match
    typical = fp_def.get('typical_pages', (1, 50))
    pages_in_range = typical[0] <= num_pages <= typical[1]
    max_pages = fp_def.get('max_pages')
    min_pages = fp_def.get('min_pages')
    if max_pages and num_pages > max_pages * 2:
        return 0.0  # Way too many pages
    if min_pages and num_pages < min_pages:
        return 0.0  # Too few pages

    # Orientation match
    expected_orient = fp_def.get('orientation')
    orient_match = True
    if expected_orient == 'landscape' and not is_landscape:
        orient_match = False

    # Producer match
    producer_hints = fp_def.get('producer_hints', [])
    metadata_match = any(h.lower() in producer or h.lower() in creator for h in producer_hints)

    # Calculate final score
    score = tier2_confidence(
        keyword_hits=keyword_hits,
        structural_match=profile_match and pages_in_range,
        metadata_match=metadata_match,
    )

    # Penalties
    if not orient_match:
        score -= 0.1
    if not pages_in_range:
        score -= 0.05

    # Boost for high keyword ratio
    if kw_ratio > 0.4:
        score += 0.05

    return max(0.0, min(score, 0.95))


def _read_excel_text(rel_path: str, abs_path: Path) -> tuple[str | None, bool]:
    """Read text from an Excel file (.xls or .xlsx). Returns (all_text, has_depth_pattern)."""
    suffix = abs_path.suffix.lower()
    all_text = ""
    has_depth_pattern = False

    if suffix == '.xlsx':
        try:
            import openpyxl
        except ImportError:
            logger.warning("openpyxl not available for .xlsx fingerprinting")
            return None, False
        try:
            wb = openpyxl.load_workbook(str(abs_path), read_only=True, data_only=True)
        except Exception as e:
            logger.warning("Cannot open Excel %s: %s", rel_path, e)
            return None, False
        try:
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                all_text += sheet_name.lower() + " "
                for row_idx, row in enumerate(ws.iter_rows(max_row=20, max_col=20, values_only=True)):
                    if row_idx >= 20:
                        break
                    for val in row:
                        if isinstance(val, str):
                            all_text += val.lower() + " "
                        elif isinstance(val, (int, float)):
                            if abs(val % 0.20) < 0.01 and 0 < val < 20:
                                has_depth_pattern = True
            wb.close()
        except Exception as e:
            logger.warning("Tier 2 Excel analysis failed for %s: %s", rel_path, e)
            return None, False
    else:
        # .xls via xlrd
        try:
            import xlrd
        except ImportError:
            logger.warning("xlrd not available for .xls fingerprinting")
            return None, False
        try:
            wb = xlrd.open_workbook(str(abs_path), on_demand=True)
        except Exception as e:
            logger.warning("Cannot open Excel %s: %s", rel_path, e)
            return None, False
        try:
            for sheet_name in wb.sheet_names():
                sheet = wb.sheet_by_name(sheet_name)
                all_text += sheet_name.lower() + " "
                for row_idx in range(min(sheet.nrows, 20)):
                    for col_idx in range(min(sheet.ncols, 20)):
                        try:
                            val = sheet.cell_value(row_idx, col_idx)
                            if isinstance(val, str):
                                all_text += val.lower() + " "
                            elif isinstance(val, (int, float)):
                                if abs(val % 0.20) < 0.01 and 0 < val < 20:
                                    has_depth_pattern = True
                        except Exception:
                            pass
            wb.release_resources()
        except Exception as e:
            logger.warning("Tier 2 Excel analysis failed for %s: %s", rel_path, e)
            return None, False

    return all_text, has_depth_pattern


def _fingerprint_excel(rel_path: str, abs_path: Path) -> FileClassification | None:
    """Analyze an Excel file's content to determine its role.

    Supports both .xls (xlrd) and .xlsx (openpyxl).
    """
    all_text, has_depth_pattern = _read_excel_text(rel_path, abs_path)
    if all_text is None:
        return None

    try:
        # Score DPSH Excel
        dpsh_score = 0.0
        dpsh_def = ROLE_FINGERPRINTS.get('dpsh_excel', {})
        dpsh_keywords = dpsh_def.get('keywords_excel', [])
        dpsh_hits = sum(1 for kw in dpsh_keywords if kw.lower() in all_text)

        if dpsh_hits >= 2:
            dpsh_score = tier2_confidence(keyword_hits=dpsh_hits)
            if has_depth_pattern:
                dpsh_score += 0.1  # Strong indicator: 0.20m depth intervals

        # Score lab order
        lab_score = 0.0
        lab_def = ROLE_FINGERPRINTS.get('lab_order', {})
        lab_keywords = lab_def.get('keywords_excel', [])
        lab_hits = sum(1 for kw in lab_keywords if kw.lower() in all_text)

        if lab_hits >= 2:
            lab_score = tier2_confidence(keyword_hits=lab_hits)

        # Return best match
        best_role: str | None = None
        best_score = 0.0
        alternates: list[dict] = []

        if dpsh_score > 0:
            alternates.append({'role': 'dpsh_excel', 'confidence': round(dpsh_score, 3)})
            if dpsh_score > best_score:
                best_score = dpsh_score
                best_role = 'dpsh_excel'

        if lab_score > 0:
            alternates.append({'role': 'lab_order', 'confidence': round(lab_score, 3)})
            if lab_score > best_score:
                best_score = lab_score
                best_role = 'lab_order'

        if best_role and best_score >= 0.3:
            from .confidence import classify_category
            return FileClassification(
                file_path=rel_path,
                role=best_role,
                confidence=best_score,
                tier=ClassificationTier.FINGERPRINT,
                alternate_roles=[a for a in alternates if a['role'] != best_role],
                fingerprint_data={'text_sample': all_text[:200], 'has_depth_pattern': has_depth_pattern},
                category=classify_category(best_score, best_role),
            )

    except Exception as e:
        logger.warning(f"Tier 2 Excel analysis failed for {rel_path}: {e}")

    return None


# ──────────────────────────────────────────────────────────────
# Image fingerprinting (Pillow)
# ──────────────────────────────────────────────────────────────

def _fingerprint_image(rel_path: str, abs_path: Path) -> FileClassification | None:
    """Analyze an image to determine if it's a document/plan photo or a field photo.

    Uses Pillow for cheap structural analysis (zero API cost):
    - Resolution: tiny images are thumbnails → informative
    - Color analysis: documents are mostly white/light; field photos are colorful
    - EXIF: phone camera → field photo; no EXIF → screenshot/scan
    - Aspect ratio: A4-like suggests scanned document

    Returns classification or None (leaves for Tier 3 vision).
    """
    try:
        from PIL import Image
        from PIL.ExifTags import Base as ExifBase
    except ImportError:
        logger.warning("Pillow not available for Tier 2 image fingerprinting")
        return None

    try:
        img = Image.open(str(abs_path))
        w, h = img.size
        file_size_kb = abs_path.stat().st_size // 1024

        # ── 1. Tiny images → informative (thumbnail/icon)
        total_pixels = w * h
        if total_pixels < 160_000:  # ~400x400
            return FileClassification(
                file_path=rel_path,
                role=None,
                confidence=0.9,
                tier=ClassificationTier.FINGERPRINT,
                category="informative",
                summary=f"thumbnail ({w}x{h})",
                fingerprint_data={'width': w, 'height': h, 'reason': 'too_small'},
            )

        # ── 2. EXIF analysis
        exif = img.getexif()
        has_camera_exif = bool(exif.get(ExifBase.Make) or exif.get(ExifBase.Model))

        # ── 3. Color analysis (sample pixels for white-dominance + saturation)
        lightness_ratio, saturation_ratio = _compute_color_profile(img)

        # ── 4. Build fingerprint data
        aspect_ratio = max(w, h) / min(w, h) if min(w, h) > 0 else 1.0
        is_a4_like = 1.3 <= aspect_ratio <= 1.55  # A4 = 1.414
        is_large = total_pixels >= 500_000  # ~700x700 or bigger

        fingerprint: dict = {
            'width': w,
            'height': h,
            'aspect_ratio': round(aspect_ratio, 2),
            'file_size_kb': file_size_kb,
            'has_camera_exif': has_camera_exif,
            'lightness_ratio': round(lightness_ratio, 2),
            'saturation_ratio': round(saturation_ratio, 2),
            'is_a4_like': is_a4_like,
        }

        img.close()

        # ── 5. Classify using a scoring approach
        # Documents (plans, field sheets, forms) are:
        #   - Mostly white/light (>50%) OR low saturation (technical drawing)
        #   - Often A4-like aspect ratio
        # Photos of documents on a surface may have dark borders but the
        # center is still light → we use a lenient threshold.
        # Field photos: colorful, high saturation, varied lightness.

        is_document_like = (
            lightness_ratio >= 0.50  # Clean scan/screenshot of document
            or (lightness_ratio >= 0.30 and saturation_ratio < 0.15)  # Desaturated = technical drawing
        )

        if is_document_like:
            summary = "document_photo" if lightness_ratio >= 0.60 else "possible_document"
            return FileClassification(
                file_path=rel_path,
                role=None,
                confidence=0.7 if lightness_ratio >= 0.60 else 0.55,
                tier=ClassificationTier.FINGERPRINT,
                category="needs_vision",
                summary=summary,
                fingerprint_data=fingerprint,
            )

        # Field photo with camera EXIF → informative
        if has_camera_exif:
            return FileClassification(
                file_path=rel_path,
                role=None,
                confidence=0.8,
                tier=ClassificationTier.FINGERPRINT,
                category="informative",
                summary="field_photo_exif",
                fingerprint_data=fingerprint,
            )

        # No EXIF, not document-like, but reasonably large → needs_vision
        # Could be a screenshot (Google Maps, geological map, site overview)
        if is_large:
            return FileClassification(
                file_path=rel_path,
                role=None,
                confidence=0.5,
                tier=ClassificationTier.FINGERPRINT,
                category="needs_vision",
                summary="screenshot_or_map",
                fingerprint_data=fingerprint,
            )

        # Fallback: unclassified image with fingerprint data
        return FileClassification(
            file_path=rel_path,
            role=None,
            confidence=0.0,
            tier=ClassificationTier.FINGERPRINT,
            category="unknown",
            summary=f"image {w}x{h}, light={lightness_ratio:.0%}",
            fingerprint_data=fingerprint,
        )

    except Exception as e:
        logger.warning(f"Tier 2 image analysis failed for {rel_path}: {e}")

    return None


def _compute_color_profile(img) -> tuple[float, float]:
    """Compute lightness ratio and saturation ratio of an image.

    Returns (lightness_ratio, saturation_ratio) both in 0.0-1.0.

    - lightness_ratio: fraction of pixels that are 'light' (white/cream/gray)
      Documents > 0.5, field photos < 0.3
    - saturation_ratio: fraction of pixels that are 'colorful' (not gray)
      Field photos > 0.3, technical drawings < 0.15
    """
    try:
        if img.mode != 'RGB':
            img = img.convert('RGB')

        # Resize to small thumbnail for fast analysis (~1ms)
        thumb = img.resize((100, 100))
        _getdata = getattr(thumb, 'get_flattened_data', None) or thumb.getdata
        pixels = list(_getdata())
        n = len(pixels)
        if n == 0:
            return 0.0, 0.0

        light_threshold = 180
        sat_threshold = 40  # Min difference between max and min channel

        light_count = 0
        saturated_count = 0

        for r, g, b in pixels:
            if r > light_threshold and g > light_threshold and b > light_threshold:
                light_count += 1
            chan_range = max(r, g, b) - min(r, g, b)
            if chan_range > sat_threshold:
                saturated_count += 1

        return light_count / n, saturated_count / n

    except Exception:
        return 0.0, 0.5
