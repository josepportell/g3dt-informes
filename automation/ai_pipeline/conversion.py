"""Stage 3: conversion — produce LLM-ready artifacts from every useful file.

Consumes the Stage 2 typology and dispatches each file to a converter based on
its `conversion_strategy`. Outputs land in `{project}/validation/ai_pipeline/converted/{source_stem}/`.

Design doc: docs/ARQUITECTURA-AI-PIPELINE.md §6
"""

from __future__ import annotations

import csv
import hashlib
import io
import logging
import os
import re
import shutil
import subprocess
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Callable

import yaml
from pydantic import BaseModel, Field

from .typology import FileClass, ProjectTypology, classify_project

logger = logging.getLogger(__name__)

# ─── Configuration ─────────────────────────────────────────────────────

_CONVERTED_ROOT = Path("validation") / "ai_pipeline" / "converted"
_PDF_RENDER_DPI = 200
_LIBREOFFICE_TIMEOUT_S = 60


# ─── Models ────────────────────────────────────────────────────────────


class ConvertedArtifact(BaseModel):
    """One artifact produced by Stage 3. Path is relative to project root."""

    path: str                       # relative POSIX
    format: str                     # "md", "csv", "png", "jpg", "passthrough"
    source_path: str                # the FileClass that produced this artifact
    source_chain: list[str] = Field(default_factory=list)
    strategy_used: str
    page: int | None = None         # 1-based for PDFs
    sheet: str | None = None        # sheet name for Excel
    bytes_written: int = 0
    skipped: bool = False
    skip_reason: str = ""


class ProjectConversion(BaseModel):
    """Stage 3 output. All artifacts + per-source grouping + eva summary."""

    project_path: str
    converted_at: str
    artifacts: list[ConvertedArtifact] = Field(default_factory=list)
    per_source: dict[str, list[str]] = Field(default_factory=dict)
    total_bytes: int = 0
    warnings: list[str] = Field(default_factory=list)
    eva_summary: list[str] = Field(default_factory=list)

    # Query helpers
    def artifacts_of(self, source_path: str) -> list[ConvertedArtifact]:
        return [a for a in self.artifacts if a.source_path == source_path]

    def artifacts_by_format(self, fmt: str) -> list[ConvertedArtifact]:
        return [a for a in self.artifacts if a.format == fmt]

    def artifacts_by_strategy(self, strategy: str) -> list[ConvertedArtifact]:
        return [a for a in self.artifacts if a.strategy_used == strategy]


# ─── Shared utilities ──────────────────────────────────────────────────


def _slugify(s: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", s).strip("._-")
    return cleaned or "unnamed"


def _out_dir(pp: Path, fc: FileClass) -> Path:
    stem = _slugify(Path(fc.path).stem)
    d = pp / _CONVERTED_ROOT / stem
    d.mkdir(parents=True, exist_ok=True)
    return d


def _rel(pp: Path, p: Path) -> str:
    return p.relative_to(pp).as_posix()


def _write_dedup(out_path: Path, data: bytes) -> int:
    """Write bytes, skipping disk I/O when content matches SHA256 of existing file.

    Always returns len(data) — the size on disk, which is independent of whether
    this call actually wrote. During dev we re-run Stage 3 freely; dedup keeps
    the cost low when nothing has changed.
    """
    out_path.parent.mkdir(parents=True, exist_ok=True)
    if out_path.is_file():
        existing = hashlib.sha256(out_path.read_bytes()).hexdigest()
        incoming = hashlib.sha256(data).hexdigest()
        if existing == incoming:
            return len(data)
    out_path.write_bytes(data)
    return len(data)


def _images_manifest_md(children: list[FileClass], manifest_path: Path, pp: Path) -> str:
    """Build an `images.md` that references extracted child images.

    Each ref is a relative path from the manifest file's directory to the image.
    """
    if not children:
        return ""
    lines = ["# Imatges embegudes\n"]
    for c in children:
        img_abs = pp / c.path
        rel = os.path.relpath(str(img_abs), str(manifest_path.parent))
        lines.append(f"![{Path(c.path).name}]({rel})")
        lines.append("")
    return "\n".join(lines)


# ─── Converters ────────────────────────────────────────────────────────


def _convert_pdf_text(pp: Path, fc: FileClass, typ: ProjectTypology) -> list[ConvertedArtifact]:
    import pymupdf4llm

    out = _out_dir(pp, fc)
    abs_path = pp / fc.path
    pages = pymupdf4llm.to_markdown(
        str(abs_path),
        page_chunks=True,
        write_images=False,
        show_progress=False,
    )
    artifacts: list[ConvertedArtifact] = []
    for i, page_data in enumerate(pages):
        md_text = page_data.get("text", "") if isinstance(page_data, dict) else str(page_data)
        out_path = out / f"page_{i+1:03d}.md"
        written = _write_dedup(out_path, md_text.encode("utf-8"))
        artifacts.append(ConvertedArtifact(
            path=_rel(pp, out_path),
            format="md",
            source_path=fc.path,
            source_chain=list(fc.source_chain),
            strategy_used=fc.conversion_strategy,
            page=i + 1,
            bytes_written=written,
        ))
    return artifacts


def _convert_pdf_mixed(pp: Path, fc: FileClass, typ: ProjectTypology) -> list[ConvertedArtifact]:
    # Per-page markdown + a separate images.md with refs to Stage 2 extractions
    artifacts = _convert_pdf_text(pp, fc, typ)
    children = typ.children_of(fc.path)
    if children:
        out = _out_dir(pp, fc)
        images_path = out / "images.md"
        md = _images_manifest_md(children, images_path, pp)
        written = _write_dedup(images_path, md.encode("utf-8"))
        artifacts.append(ConvertedArtifact(
            path=_rel(pp, images_path),
            format="md",
            source_path=fc.path,
            source_chain=list(fc.source_chain),
            strategy_used=fc.conversion_strategy,
            bytes_written=written,
        ))
    return artifacts


def _convert_pdf_scanned(pp: Path, fc: FileClass, typ: ProjectTypology) -> list[ConvertedArtifact]:
    import fitz

    out = _out_dir(pp, fc)
    abs_path = pp / fc.path
    doc = fitz.open(str(abs_path))
    artifacts: list[ConvertedArtifact] = []
    try:
        for i in range(doc.page_count):
            page = doc[i]
            pix = page.get_pixmap(dpi=_PDF_RENDER_DPI)
            data = pix.tobytes("png")
            out_path = out / f"page_{i+1:03d}.png"
            written = _write_dedup(out_path, data)
            artifacts.append(ConvertedArtifact(
                path=_rel(pp, out_path),
                format="png",
                source_path=fc.path,
                source_chain=list(fc.source_chain),
                strategy_used=fc.conversion_strategy,
                page=i + 1,
                bytes_written=written,
            ))
    finally:
        doc.close()
    return artifacts


def _convert_docx(pp: Path, fc: FileClass, typ: ProjectTypology) -> list[ConvertedArtifact]:
    import pypandoc

    out = _out_dir(pp, fc)
    abs_path = pp / fc.path
    try:
        md_text = pypandoc.convert_file(str(abs_path), to="gfm")
    except Exception as e:
        return [ConvertedArtifact(
            path="", format="", source_path=fc.path,
            source_chain=list(fc.source_chain),
            strategy_used=fc.conversion_strategy,
            skipped=True, skip_reason=f"pandoc error: {e}",
        )]

    out_path = out / "body.md"
    written = _write_dedup(out_path, md_text.encode("utf-8"))
    artifacts = [ConvertedArtifact(
        path=_rel(pp, out_path),
        format="md",
        source_path=fc.path,
        source_chain=list(fc.source_chain),
        strategy_used=fc.conversion_strategy,
        bytes_written=written,
    )]

    children = typ.children_of(fc.path)
    if children:
        images_path = out / "images.md"
        md = _images_manifest_md(children, images_path, pp)
        img_written = _write_dedup(images_path, md.encode("utf-8"))
        artifacts.append(ConvertedArtifact(
            path=_rel(pp, images_path),
            format="md",
            source_path=fc.path,
            source_chain=list(fc.source_chain),
            strategy_used=fc.conversion_strategy,
            bytes_written=img_written,
        ))
    return artifacts


def _convert_doc_legacy(pp: Path, fc: FileClass, typ: ProjectTypology) -> list[ConvertedArtifact]:
    """Convert legacy .doc via LibreOffice headless → .docx → pypandoc."""
    libreoffice = shutil.which("libreoffice") or shutil.which("soffice")
    if not libreoffice:
        return [ConvertedArtifact(
            path="", format="", source_path=fc.path,
            source_chain=list(fc.source_chain),
            strategy_used=fc.conversion_strategy,
            skipped=True,
            skip_reason="legacy .doc; LibreOffice binary not found on PATH",
        )]

    import pypandoc

    abs_path = pp / fc.path
    with tempfile.TemporaryDirectory(prefix="g3dt_doc_") as tmp:
        tmp_path = Path(tmp)
        try:
            subprocess.run(
                [libreoffice, "--headless", "--convert-to", "docx",
                 "--outdir", str(tmp_path), str(abs_path)],
                check=True, capture_output=True,
                timeout=_LIBREOFFICE_TIMEOUT_S,
            )
        except subprocess.TimeoutExpired:
            return [ConvertedArtifact(
                path="", format="", source_path=fc.path,
                source_chain=list(fc.source_chain),
                strategy_used=fc.conversion_strategy,
                skipped=True, skip_reason="LibreOffice timeout",
            )]
        except subprocess.CalledProcessError as e:
            return [ConvertedArtifact(
                path="", format="", source_path=fc.path,
                source_chain=list(fc.source_chain),
                strategy_used=fc.conversion_strategy,
                skipped=True,
                skip_reason=f"LibreOffice error: {e.stderr.decode('utf-8', errors='replace')[:200]}",
            )]

        tmp_docx = tmp_path / f"{abs_path.stem}.docx"
        if not tmp_docx.is_file():
            # LibreOffice sometimes sanitizes filenames — pick whatever .docx appeared
            found = list(tmp_path.glob("*.docx"))
            if not found:
                return [ConvertedArtifact(
                    path="", format="", source_path=fc.path,
                    source_chain=list(fc.source_chain),
                    strategy_used=fc.conversion_strategy,
                    skipped=True, skip_reason="LibreOffice produced no .docx",
                )]
            tmp_docx = found[0]

        try:
            md_text = pypandoc.convert_file(str(tmp_docx), to="gfm")
        except Exception as e:
            return [ConvertedArtifact(
                path="", format="", source_path=fc.path,
                source_chain=list(fc.source_chain),
                strategy_used=fc.conversion_strategy,
                skipped=True, skip_reason=f"pandoc error: {e}",
            )]

    out = _out_dir(pp, fc)
    out_path = out / "body.md"
    written = _write_dedup(out_path, md_text.encode("utf-8"))
    return [ConvertedArtifact(
        path=_rel(pp, out_path),
        format="md",
        source_path=fc.path,
        source_chain=list(fc.source_chain),
        strategy_used=fc.conversion_strategy,
        bytes_written=written,
    )]


def _sheet_to_csv(rows: list[list]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf, quoting=csv.QUOTE_MINIMAL)
    for row in rows:
        writer.writerow(["" if v is None else str(v) for v in row])
    return buf.getvalue().encode("utf-8")


def _convert_excel(pp: Path, fc: FileClass, typ: ProjectTypology) -> list[ConvertedArtifact]:
    out = _out_dir(pp, fc)
    abs_path = pp / fc.path
    artifacts: list[ConvertedArtifact] = []

    if fc.format == "xlsx":
        import openpyxl
        try:
            wb = openpyxl.load_workbook(str(abs_path), data_only=True, read_only=True)
        except Exception as e:
            return [ConvertedArtifact(
                path="", format="", source_path=fc.path,
                source_chain=list(fc.source_chain),
                strategy_used=fc.conversion_strategy,
                skipped=True, skip_reason=f"openpyxl error: {e}",
            )]
        try:
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                rows = [list(r) for r in ws.iter_rows(values_only=True)]
                # Trim trailing empty rows for brevity
                while rows and not any(cell is not None and str(cell).strip() for cell in rows[-1]):
                    rows.pop()
                if not rows:
                    continue
                data = _sheet_to_csv(rows)
                out_path = out / f"sheet_{_slugify(sheet_name)}.csv"
                written = _write_dedup(out_path, data)
                artifacts.append(ConvertedArtifact(
                    path=_rel(pp, out_path),
                    format="csv",
                    source_path=fc.path,
                    source_chain=list(fc.source_chain),
                    strategy_used=fc.conversion_strategy,
                    sheet=sheet_name,
                    bytes_written=written,
                ))
        finally:
            wb.close()
    elif fc.format == "xls":
        try:
            import xlrd
        except ImportError:
            return [ConvertedArtifact(
                path="", format="", source_path=fc.path,
                source_chain=list(fc.source_chain),
                strategy_used=fc.conversion_strategy,
                skipped=True, skip_reason="xlrd not installed (needed for legacy .xls)",
            )]
        try:
            book = xlrd.open_workbook(str(abs_path))
        except Exception as e:
            return [ConvertedArtifact(
                path="", format="", source_path=fc.path,
                source_chain=list(fc.source_chain),
                strategy_used=fc.conversion_strategy,
                skipped=True, skip_reason=f"xlrd error: {e}",
            )]
        for sheet_idx in range(book.nsheets):
            sh = book.sheet_by_index(sheet_idx)
            rows = []
            for r in range(sh.nrows):
                rows.append([sh.cell_value(r, c) for c in range(sh.ncols)])
            while rows and not any(cell != "" and cell is not None for cell in rows[-1]):
                rows.pop()
            if not rows:
                continue
            data = _sheet_to_csv(rows)
            out_path = out / f"sheet_{_slugify(sh.name)}.csv"
            written = _write_dedup(out_path, data)
            artifacts.append(ConvertedArtifact(
                path=_rel(pp, out_path),
                format="csv",
                source_path=fc.path,
                source_chain=list(fc.source_chain),
                strategy_used=fc.conversion_strategy,
                sheet=sh.name,
                bytes_written=written,
            ))

    # Append images manifest for spreadsheet_mixed
    if fc.conversion_strategy == "excel_per_sheet_to_csv_plus_images":
        children = typ.children_of(fc.path)
        if children:
            images_path = out / "images.md"
            md = _images_manifest_md(children, images_path, pp)
            img_written = _write_dedup(images_path, md.encode("utf-8"))
            artifacts.append(ConvertedArtifact(
                path=_rel(pp, images_path),
                format="md",
                source_path=fc.path,
                source_chain=list(fc.source_chain),
                strategy_used=fc.conversion_strategy,
                bytes_written=img_written,
            ))

    return artifacts


def _convert_msg(pp: Path, fc: FileClass, typ: ProjectTypology) -> list[ConvertedArtifact]:
    try:
        import extract_msg as _em
    except ImportError:
        return [ConvertedArtifact(
            path="", format="", source_path=fc.path,
            source_chain=list(fc.source_chain),
            strategy_used=fc.conversion_strategy,
            skipped=True, skip_reason="extract_msg not installed",
        )]

    abs_path = pp / fc.path
    try:
        m = _em.Message(str(abs_path))
    except Exception as e:
        return [ConvertedArtifact(
            path="", format="", source_path=fc.path,
            source_chain=list(fc.source_chain),
            strategy_used=fc.conversion_strategy,
            skipped=True, skip_reason=f"extract_msg error: {e}",
        )]

    try:
        body = m.body or ""
        # YAML frontmatter with email metadata — use yaml.safe_dump so values
        # with colons, newlines, or unicode accents are properly escaped.
        def _yaml_value(v) -> str:
            if v is None:
                return ""
            return str(v).replace("\n", " ").strip()

        frontmatter = {
            "from": _yaml_value(m.sender),
            "to": _yaml_value(m.to),
            "subject": _yaml_value(m.subject),
            "date": _yaml_value(m.date),
        }
        yaml_block = yaml.safe_dump(
            frontmatter,
            default_flow_style=False,
            allow_unicode=True,
            sort_keys=False,
        )
        content = f"---\n{yaml_block}---\n{body}"
    finally:
        try:
            m.close()
        except Exception:
            pass

    out = _out_dir(pp, fc)
    out_path = out / "body.md"
    written = _write_dedup(out_path, content.encode("utf-8"))
    return [ConvertedArtifact(
        path=_rel(pp, out_path),
        format="md",
        source_path=fc.path,
        source_chain=list(fc.source_chain),
        strategy_used=fc.conversion_strategy,
        bytes_written=written,
    )]


def _passthrough(pp: Path, fc: FileClass, typ: ProjectTypology) -> list[ConvertedArtifact]:
    """For image/text passthrough: reference the original file, no conversion.

    Extracted images (`parent_path` set) inherit the parent's `source_path` so
    Stage 4 groups them with the parent document rather than issuing a separate
    LLM call per image (D1#2).
    """
    abs_path = pp / fc.path
    size = abs_path.stat().st_size if abs_path.is_file() else 0
    source_path = fc.path
    if fc.parent_path and any(other.path == fc.parent_path for other in typ.files):
        source_path = fc.parent_path
    return [ConvertedArtifact(
        path=fc.path,                    # points at the original file
        format="passthrough",
        source_path=source_path,
        source_chain=list(fc.source_chain),
        strategy_used=fc.conversion_strategy,
        bytes_written=size,
    )]


# ─── Dispatch ──────────────────────────────────────────────────────────


_DISPATCH: dict[str, Callable[[Path, FileClass, ProjectTypology], list[ConvertedArtifact]]] = {
    "pdf_to_markdown": _convert_pdf_text,
    "pdf_to_markdown_plus_images": _convert_pdf_mixed,
    "pdf_pages_to_images": _convert_pdf_scanned,
    "docx_to_markdown_plus_media": _convert_docx,  # .docx path; .doc handled below
    "excel_per_sheet_to_csv": _convert_excel,
    "excel_per_sheet_to_csv_plus_images": _convert_excel,
    "msg_body_to_markdown": _convert_msg,
    "image_passthrough": _passthrough,
    "text_passthrough": _passthrough,
}


def _dispatch_for(fc: FileClass):
    """Pick the right converter, handling the .doc vs .docx format split."""
    if fc.conversion_strategy == "docx_to_markdown_plus_media" and fc.format == "doc":
        return _convert_doc_legacy
    return _DISPATCH.get(fc.conversion_strategy)


# ─── Eva summary ───────────────────────────────────────────────────────


def _build_eva_summary(
    artifacts: list[ConvertedArtifact],
    typ: ProjectTypology,
    warnings: list[str],
) -> list[str]:
    # Unique source files that actually produced artifacts — respects strategy filter
    useful_sources = len({a.source_path for a in artifacts if not a.skipped and a.path})

    by_fmt: dict[str, int] = {}
    for a in artifacts:
        if a.skipped:
            continue
        by_fmt[a.format] = by_fmt.get(a.format, 0) + 1

    md_pages = by_fmt.get("md", 0)
    csv_sheets = by_fmt.get("csv", 0)
    png_pages = by_fmt.get("png", 0)
    passthrough = by_fmt.get("passthrough", 0)
    skipped = sum(1 for a in artifacts if a.skipped)

    total = sum(1 for a in artifacts if not a.skipped)
    lines = [
        f"Hem convertit {useful_sources} fitxers útils a {total} artefactes llegibles per la IA:",
    ]

    if md_pages:
        lines.append(f"• {md_pages} pàgines/documents en markdown.")
    if csv_sheets:
        lines.append(f"• {csv_sheets} fulls de càlcul en CSV.")
    if png_pages:
        lines.append(f"• {png_pages} pàgines de PDFs escanejats com a imatges (per lectura visual a Fase 4).")
    if passthrough:
        lines.append(f"• {passthrough} imatges i fitxers de text referenciats com estan (no cal conversió).")
    if skipped:
        lines.append(f"• {skipped} fitxers no s'han pogut convertir — vegeu advertències.")

    return lines


# ─── Public API ────────────────────────────────────────────────────────


def convert_project(
    project_path: Path | str,
    *,
    typology: ProjectTypology | None = None,
    strategies: set[str] | None = None,
) -> ProjectConversion:
    """Convert every useful file in the project to LLM-ready artifacts.

    Args:
        project_path: Absolute or relative path to the project folder.
        typology: Optionally reuse an existing Stage 2 typology. If None, built fresh.
        strategies: Optional subset of `conversion_strategy` values to process.
            Useful for dev: `strategies={"pdf_to_markdown"}` runs only that converter.

    Returns:
        ProjectConversion with artifacts + manifest + eva_summary.
    """
    pp = Path(project_path).resolve()
    if not pp.is_dir():
        raise ValueError(f"Not a directory: {pp}")

    typ = typology if typology is not None else classify_project(pp)

    artifacts: list[ConvertedArtifact] = []
    warnings: list[str] = []

    for fc in typ.files:
        if not fc.useful or fc.conversion_strategy == "skip":
            continue
        if strategies and fc.conversion_strategy not in strategies:
            continue

        converter = _dispatch_for(fc)
        if converter is None:
            warnings.append(f"no converter for strategy: {fc.conversion_strategy} ({fc.path})")
            continue

        try:
            new = converter(pp, fc, typ)
            artifacts.extend(new)
        except Exception as e:
            logger.warning("Conversion failed for %s: %s", fc.path, e, exc_info=True)
            artifacts.append(ConvertedArtifact(
                path="",
                format="",
                source_path=fc.path,
                source_chain=list(fc.source_chain),
                strategy_used=fc.conversion_strategy,
                skipped=True,
                skip_reason=f"converter error: {e}",
            ))

    # Aggregate
    per_source: dict[str, list[str]] = {}
    total_bytes = 0
    for a in artifacts:
        if a.skipped or not a.path:
            continue
        per_source.setdefault(a.source_path, []).append(a.path)
        total_bytes += a.bytes_written

    eva_summary = _build_eva_summary(artifacts, typ, warnings)

    return ProjectConversion(
        project_path=str(pp),
        converted_at=datetime.now(timezone.utc).isoformat(timespec="seconds"),
        artifacts=artifacts,
        per_source=per_source,
        total_bytes=total_bytes,
        warnings=warnings,
        eva_summary=eva_summary,
    )


def save_conversion(conv: ProjectConversion, project_path: Path | str) -> Path:
    """Write manifest to {project}/validation/ai_conversion.json."""
    pp = Path(project_path).resolve()
    out_dir = pp / "validation"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / "ai_conversion.json"
    out_path.write_text(conv.model_dump_json(indent=2), encoding="utf-8")
    return out_path


def load_conversion(project_path: Path | str) -> ProjectConversion | None:
    pp = Path(project_path).resolve()
    p = pp / "validation" / "ai_conversion.json"
    if not p.is_file():
        return None
    return ProjectConversion.model_validate_json(p.read_text(encoding="utf-8"))
