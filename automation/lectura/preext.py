"""Experiment de pre-extracció determinista (via A, lectura headless).

Sonda `docs/wizard-headless/mesures/probes/2026-08-24-tall-stream-json/turns.md`:
de 15 usos d'eina d'una crida `--only` típica, 5 són el model construint-se
l'eina (`fitz` per comptar pàgines, extreure text, renderitzar sencer/meitats/
retalls) i 2 més són cerimònia d'escriptura (`os.replace` manual, rellegir).
Aquest mòdul fa aquesta feina ABANS que el model hi entri: per a cada document
de l'inventari amb `route == "claude"`, escriu a disc exactament el que el
model hauria construït (text per pàgina, PNG de pàgina sencera i meitats,
CSV/cel·les d'Excel, cos+adjunts de correu, imatge reduïda), perquè el skill
només hagi de `Read` i escriure el JSON amb `scripts/write_doc_json.py`.

Determinisme: cap accés de xarxa, cap escriptura dins `project_path` (només a
`preext_root`). Idempotent per `source_md5`: una segona crida sobre el mateix
document no torna a escriure res.

API pública: `preext_document()` (un sol document) i `augment_inventory()`
(tot l'inventari — afegeix la clau `preext` a cada entrada `route == "claude"`
i reescriu `_inventory.json`).

Va SEMPRE darrere del flag `G3DT_LECTURA_PREEXT` (`automation/lectura/runner.py`):
apagat, aquest mòdul ni tan sols s'importa.
"""

from __future__ import annotations

import contextlib
import csv
import hashlib
import json
import logging
import os
import re
import tempfile
import time
from datetime import date, datetime
from pathlib import Path, PurePosixPath
from typing import Any

from automation.lectura.runner import safe_doc_name

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Constants d'extracció (brief de l'experiment)
# ---------------------------------------------------------------------------

_MAX_PAGES = 12
_WHOLE_PAGE_DPI = 150
_HALF_DPI = 220
_CLIP_MAIN_FRACTION = 0.55  # 0-55 %
_CLIP_OTHER_START = 0.45  # 45-100 % (10 % de solapament)
_MIN_TEXT_ALNUM_CHARS = 20

_IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png"}
_IMAGE_MAX_LONG_SIDE = 1600
_MSG_ATTACHMENT_MIN_BYTES = 25 * 1024  # signatures ≤ 25 KB, descartades
_MSG_ATTACHMENT_RECURSE_EXTENSIONS = {".pdf", ".xls", ".xlsx"} | _IMAGE_EXTENSIONS

_ATTACH_NAME_RE = re.compile(r"[^A-Za-z0-9]+")


# ---------------------------------------------------------------------------
# Utilitats d'I/O
# ---------------------------------------------------------------------------


def _write_json_atomic(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp_name = tempfile.mkstemp(dir=str(path.parent), prefix=path.name + ".", suffix=".tmp")
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as fh:
            json.dump(payload, fh, ensure_ascii=False, indent=1, default=str)
        os.replace(tmp_name, path)
    except Exception:
        if os.path.exists(tmp_name):
            os.remove(tmp_name)
        raise


def _try_load_json(path: Path) -> Any:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return None


def _rel(preext_root: Path, path: Path) -> str:
    """Path relatiu a `preext_root`, sempre amb separador '/' (multiplataforma)."""
    return Path(path).resolve().relative_to(preext_root.resolve()).as_posix()


def _safe_attachment_name(name: str) -> str:
    """Sanejada d'un nom d'adjunt (mateixa regla que `safe_doc_name`, extensió intacta)."""
    p = PurePosixPath(str(name).replace("\\", "/"))
    stem = p.stem
    suffix = p.suffix.lower()
    safe_stem = _ATTACH_NAME_RE.sub("_", stem).strip("_").lower() or "adjunt"
    return f"{safe_stem}{suffix}"


def _guess_kind(suffix: str) -> str:
    suffix = suffix.lower()
    if suffix == ".pdf":
        return "pdf"
    if suffix in (".xls", ".xlsx"):
        return "excel"
    if suffix == ".msg":
        return "msg"
    if suffix in _IMAGE_EXTENSIONS:
        return "image"
    return "unsupported"


# ---------------------------------------------------------------------------
# PDF (fitz/PyMuPDF)
# ---------------------------------------------------------------------------


def _extract_pdf(source_path: Path, doc_dir: Path, preext_root: Path) -> tuple[str, list[str], dict]:
    import fitz  # PyMuPDF

    doc = fitz.open(str(source_path))
    try:
        n_pages = doc.page_count
        meta_doc = doc.metadata or {}
        files: list[str] = []
        text_ok: dict[str, bool] = {}
        page_sizes: list[list[float]] = []

        for i in range(n_pages):
            page = doc.load_page(i)
            p = i + 1
            rect = page.rect
            page_sizes.append([round(rect.width, 2), round(rect.height, 2)])

            text = page.get_text("text") or ""
            alnum_len = sum(1 for ch in text if ch.isalnum())
            text_ok[str(p)] = alnum_len >= _MIN_TEXT_ALNUM_CHARS

            txt_path = doc_dir / f"page-{p}.txt"
            txt_path.write_text(text, encoding="utf-8")
            files.append(_rel(preext_root, txt_path))

            if p > _MAX_PAGES:
                continue  # pages_truncated: només text a partir de la 13a

            png_path = doc_dir / f"page-{p}.png"
            page.get_pixmap(dpi=_WHOLE_PAGE_DPI).save(str(png_path))
            files.append(_rel(preext_root, png_path))

            width, height = rect.width, rect.height
            if width > height:
                clip_a = fitz.Rect(0, 0, width * _CLIP_MAIN_FRACTION, height)
                clip_b = fitz.Rect(width * _CLIP_OTHER_START, 0, width, height)
                name_a, name_b = "left", "right"
            else:
                clip_a = fitz.Rect(0, 0, width, height * _CLIP_MAIN_FRACTION)
                clip_b = fitz.Rect(0, height * _CLIP_OTHER_START, width, height)
                name_a, name_b = "top", "bottom"

            for clip, half_name in ((clip_a, name_a), (clip_b, name_b)):
                half_path = doc_dir / f"page-{p}-{half_name}.png"
                page.get_pixmap(dpi=_HALF_DPI, clip=clip).save(str(half_path))
                files.append(_rel(preext_root, half_path))

        extra = {
            "pages": n_pages,
            "page_sizes_pt": page_sizes,
            "producer": meta_doc.get("producer"),
            "creator": meta_doc.get("creator"),
            "creationDate": meta_doc.get("creationDate"),
            "modDate": meta_doc.get("modDate"),
            "text_ok": text_ok,
            "pages_truncated": n_pages > _MAX_PAGES,
        }
        return "pdf", files, extra
    finally:
        doc.close()


# ---------------------------------------------------------------------------
# Excel (.xls via xlrd, .xlsx via openpyxl)
# ---------------------------------------------------------------------------


def _col_letter(idx: int) -> str:
    """Índex de columna 0-based -> lletres Excel (A, B, ..., Z, AA, ...)."""
    idx += 1
    letters = ""
    while idx > 0:
        idx, rem = divmod(idx - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


def _read_xls_sheets(source_path: Path) -> list[tuple[str, list[list[Any]]]]:
    import xlrd

    book = xlrd.open_workbook(str(source_path), formatting_info=False)
    out: list[tuple[str, list[list[Any]]]] = []
    for sheet in book.sheets():
        rows: list[list[Any]] = []
        for r in range(sheet.nrows):
            row_vals = [_xls_cell_value(sheet.cell(r, c), book) for c in range(sheet.ncols)]
            rows.append(row_vals)
        out.append((sheet.name, rows))
    return out


def _xls_cell_value(cell: Any, book: Any) -> Any:
    import xlrd

    if cell.ctype == xlrd.XL_CELL_EMPTY:
        return None
    if cell.ctype == xlrd.XL_CELL_DATE:
        return xlrd.xldate_as_datetime(cell.value, book.datemode).isoformat()
    if cell.ctype == xlrd.XL_CELL_NUMBER:
        return int(cell.value) if float(cell.value).is_integer() else cell.value
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    return cell.value


def _read_xlsx_sheets(source_path: Path) -> list[tuple[str, list[list[Any]]]]:
    import openpyxl

    wb = openpyxl.load_workbook(str(source_path), data_only=True, read_only=True)
    try:
        out: list[tuple[str, list[list[Any]]]] = []
        for ws in wb.worksheets:
            rows: list[list[Any]] = []
            for row in ws.iter_rows(values_only=True):
                row_vals = []
                for v in row:
                    if isinstance(v, (datetime, date)):
                        row_vals.append(v.isoformat())
                    elif isinstance(v, float) and v.is_integer():
                        row_vals.append(int(v))
                    else:
                        row_vals.append(v)
                rows.append(row_vals)
            out.append((ws.title, rows))
        return out
    finally:
        wb.close()


def _extract_excel(source_path: Path, doc_dir: Path, preext_root: Path, suffix: str) -> tuple[str, list[str], dict]:
    sheets_data = _read_xls_sheets(source_path) if suffix == ".xls" else _read_xlsx_sheets(source_path)

    files: list[str] = []
    sheets_meta: list[dict] = []
    for i, (name, rows) in enumerate(sheets_data):
        n_cols = max((len(r) for r in rows), default=0)
        sheets_meta.append({"index": i, "name": name, "rows": len(rows), "cols": n_cols})

        csv_path = doc_dir / f"sheet-{i}.csv"
        with csv_path.open("w", encoding="utf-8", newline="") as fh:
            writer = csv.writer(fh)
            for row in rows:
                writer.writerow(["" if v is None else v for v in row])
        files.append(_rel(preext_root, csv_path))

        cells_path = doc_dir / f"sheet-{i}.cells.txt"
        lines = []
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                if val is None or val == "":
                    continue
                lines.append(f"{_col_letter(c)}{r + 1}\t{val}")
        cells_path.write_text("\n".join(lines), encoding="utf-8")
        files.append(_rel(preext_root, cells_path))

    return "excel", files, {"sheets": sheets_meta}


# ---------------------------------------------------------------------------
# Correu (.msg via extract_msg)
# ---------------------------------------------------------------------------


_HTML_TAG_RE = re.compile(r"<[^>]+>")


def _strip_html(html: str) -> str:
    return " ".join(_HTML_TAG_RE.sub(" ", html).split())


def _extract_msg(
    source_path: Path, doc_dir: Path, preext_root: Path, md5_index: dict[str, str],
) -> tuple[str, list[str], dict]:
    import extract_msg

    msg = extract_msg.Message(str(source_path))
    try:
        subject = msg.subject or ""
        sender = msg.sender or ""
        to = msg.to or ""
        date_str = str(msg.date) if msg.date else ""

        body = msg.body
        if not body:
            html = msg.htmlBody
            if isinstance(html, bytes):
                html = html.decode("utf-8", errors="replace")
            body = _strip_html(html) if html else ""

        header = f"Subject: {subject}\nFrom: {sender}\nTo: {to}\nDate: {date_str}\n\n"
        body_path = doc_dir / "body.txt"
        body_path.write_text(header + (body or ""), encoding="utf-8")
        files = [_rel(preext_root, body_path)]

        attachments_meta: list[dict] = []
        attachments_dir = doc_dir / "attachments"
        for idx, att in enumerate(msg.attachments):
            data = att.data
            if not isinstance(data, (bytes, bytearray)):
                continue  # p. ex. .msg incrustat: fora d'abast (profunditat 1)

            raw_name = att.getFilename() or att.longFilename or att.shortFilename or f"attachment_{idx}"
            size = len(data)
            ext = PurePosixPath(str(raw_name).replace("\\", "/")).suffix.lower()

            if ext in _IMAGE_EXTENSIONS and size <= _MSG_ATTACHMENT_MIN_BYTES:
                continue  # signatura, descarta (Pas 1 del skill)

            md5 = hashlib.md5(data).hexdigest()
            already = md5_index.get(md5)
            safe_name = _safe_attachment_name(raw_name)

            entry_meta = {
                "name": safe_name, "size": size, "md5": md5,
                "already_in_folder": already, "preext_dir": None,
            }
            if already:
                attachments_meta.append(entry_meta)
                continue  # ja solt a la carpeta: no dupliquis feina

            attachments_dir.mkdir(parents=True, exist_ok=True)
            if (attachments_dir / safe_name).exists():
                stem = PurePosixPath(safe_name).stem
                safe_name = f"{stem}_{idx}{ext}"
                entry_meta["name"] = safe_name

            att_path = attachments_dir / safe_name
            att_path.write_bytes(data)
            files.append(_rel(preext_root, att_path))

            if ext in _MSG_ATTACHMENT_RECURSE_EXTENSIONS:
                nested_dir = attachments_dir / f"{safe_name}.preext"
                nested_dir.mkdir(parents=True, exist_ok=True)
                try:
                    _kind, nested_files, _extra = _extract_by_kind(
                        att_path, nested_dir, safe_name, preext_root, md5_index,
                    )
                    entry_meta["preext_dir"] = _rel(preext_root, nested_dir)
                    files.extend(nested_files)
                except Exception as exc:  # pre-extracció mai fa caure la lectura
                    logger.exception("preext .msg: error al adjunt %s de %s", safe_name, source_path)
                    entry_meta["error"] = f"{type(exc).__name__}: {exc}"
            # .msg incrustat (profunditat 1): només es copia, sense recursió.

            attachments_meta.append(entry_meta)

        extra = {"subject": subject, "from": sender, "to": to, "date": date_str, "attachments": attachments_meta}
        return "msg", files, extra
    finally:
        with contextlib.suppress(Exception):
            msg.close()


# ---------------------------------------------------------------------------
# Imatge (jpg/jpeg/png via PIL)
# ---------------------------------------------------------------------------


def _extract_image(source_path: Path, doc_dir: Path, preext_root: Path) -> tuple[str, list[str], dict]:
    from PIL import Image

    original_size = source_path.stat().st_size
    with Image.open(source_path) as im:
        im.load()
        width, height = im.size
        long_side = max(width, height)
        if long_side > _IMAGE_MAX_LONG_SIDE:
            scale = _IMAGE_MAX_LONG_SIDE / long_side
            new_size = (max(1, round(width * scale)), max(1, round(height * scale)))
            im = im.convert("RGB") if im.mode not in ("RGB", "RGBA", "L") else im
            im = im.resize(new_size, Image.LANCZOS)
        out_path = doc_dir / "image.png"
        im.save(out_path, format="PNG")
        out_w, out_h = im.size

    return "image", [_rel(preext_root, out_path)], {
        "width": out_w, "height": out_h, "original_size": original_size,
    }


# ---------------------------------------------------------------------------
# No suportat (.dwg i qualsevol altra extensió)
# ---------------------------------------------------------------------------


def _extract_unsupported(source_path: Path) -> tuple[str, list[str], dict]:
    reason = f"extensió no suportada per pre-extracció: {source_path.suffix or '(cap)'}"
    return "unsupported", [], {"unsupported": True, "reason": reason}


# ---------------------------------------------------------------------------
# Dispatcher
# ---------------------------------------------------------------------------


def _extract_by_kind(
    source_path: Path, doc_dir: Path, rel_path: str, preext_root: Path, md5_index: dict[str, str],
) -> tuple[str, list[str], dict]:
    suffix = source_path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(source_path, doc_dir, preext_root)
    if suffix in (".xls", ".xlsx"):
        return _extract_excel(source_path, doc_dir, preext_root, suffix)
    if suffix == ".msg":
        return _extract_msg(source_path, doc_dir, preext_root, md5_index)
    if suffix in _IMAGE_EXTENSIONS:
        return _extract_image(source_path, doc_dir, preext_root)
    return _extract_unsupported(source_path)


# ---------------------------------------------------------------------------
# API pública
# ---------------------------------------------------------------------------


def preext_document(
    project_path: Path, entry: dict, preext_root: Path, *, md5_index: dict[str, str] | None = None,
) -> dict:
    """Pre-extreu UN document de l'inventari (idempotent per `entry["md5"]`).

    Escriu `{preext_root}/{safe_doc_name(entry["path"])}/` amb el text/imatges/
    CSV segons el tipus (Pas dispatcher `_extract_by_kind`) i `meta.json`
    (atòmic). Mai llença: qualsevol excepció durant l'extracció es captura i
    es deixa constància a `meta["error"]`, perquè la pre-extracció no faci
    mai caure la lectura headless.
    """
    project_path = Path(project_path)
    preext_root = Path(preext_root)
    rel_path = str(entry["path"])
    doc_dir = preext_root / safe_doc_name(rel_path)
    meta_path = doc_dir / "meta.json"

    cached = _try_load_json(meta_path)
    if isinstance(cached, dict) and cached.get("source_md5") == entry.get("md5"):
        return cached

    doc_dir.mkdir(parents=True, exist_ok=True)
    source_path = project_path / rel_path

    start = time.monotonic()
    try:
        kind, files, extra = _extract_by_kind(source_path, doc_dir, rel_path, preext_root, md5_index or {})
    except Exception as exc:
        logger.exception("preext_document: error processant %s", rel_path)
        kind = _guess_kind(source_path.suffix)
        files, extra = [], {"error": f"{type(exc).__name__}: {exc}"}
    elapsed_s = round(time.monotonic() - start, 3)

    meta = {
        "source_path": rel_path,
        "source_md5": entry.get("md5"),
        "kind": kind,
        "generated": datetime.now().isoformat(timespec="seconds"),
        "elapsed_s": elapsed_s,
        "dir": doc_dir.relative_to(preext_root).as_posix(),
        "files": files,
        **extra,
    }
    _write_json_atomic(meta_path, meta)
    return meta


def augment_inventory(project_path: Path, inv_path: Path, preext_root: Path | None = None) -> dict:
    """Pre-extreu tots els documents `route == "claude"` de `_inventory.json`
    i hi afegeix la clau `preext` (per entrada + bloc superior). Reescriu
    `_inventory.json` de forma atòmica. Idempotent (delega a `preext_document`).
    """
    project_path = Path(project_path)
    inv_path = Path(inv_path)
    preext_root = Path(preext_root) if preext_root is not None else inv_path.parent / "_preext"
    preext_root.mkdir(parents=True, exist_ok=True)

    inventory = json.loads(inv_path.read_text(encoding="utf-8"))
    entries = inventory.get("files", [])

    md5_index: dict[str, str] = {}
    for e in entries:
        md5_index.setdefault(e.get("md5"), e.get("path"))

    start = time.monotonic()
    n_docs = 0
    n_errors = 0
    for entry in entries:
        if entry.get("route") != "claude":
            continue
        n_docs += 1
        meta = preext_document(project_path, entry, preext_root, md5_index=md5_index)
        if meta.get("error"):
            n_errors += 1

        doc_dir = preext_root / meta["dir"]
        entry["preext"] = {
            "dir": str(doc_dir.resolve()),
            "kind": meta.get("kind"),
            "files": [str((preext_root / f).resolve()) for f in meta.get("files", [])],
            "meta": str((doc_dir / "meta.json").resolve()),
            "pages": meta.get("pages"),
            "sheets": meta.get("sheets"),
            "error": meta.get("error"),
        }

    inventory["preext"] = {
        "root": str(preext_root.resolve()),
        "n_docs": n_docs,
        "n_errors": n_errors,
        "elapsed_s": round(time.monotonic() - start, 3),
    }

    _write_json_atomic(inv_path, inventory)
    return inventory
