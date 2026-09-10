#!/usr/bin/env python3
"""Fase 8b — converteix la taula «Assaigs SPT / MA» de la plantilla en un bucle.

Per què: la plantilla tenia UNA fila de dades amb escalars (`{{ spt_test_id }}`
…), i els informes signats de l'Eva en porten més d'una quan el projecte té més
d'una mostra (Anciles: 3 files; 2 sondeigs). Amb la lectura de la via A omplint
`spt_ma_tests[]`, la fila única perdia dades reals sense avisar.

Idempotent: si la taula ja és un bucle, no fa res. Compatible amb la via B —
`report_generator` sempre emet `spt_ma_tests` (una fila construïda amb els
mateixos escalars d'abans quan no hi ha lectura), així que la sortida d'un
projecte sense via A és idèntica a la d'abans.

    python3 scripts/template_spt_ma_loop.py [templates/g3dt-jinja-template.docx]
"""

from __future__ import annotations

import copy
import sys
from pathlib import Path

from docx import Document

LOOP_OPEN = "{%tr for test in spt_ma_tests %}"
LOOP_CLOSE = "{%tr endfor %}"
DATA_CELLS = (
    "{{ test.test_id }}",
    "{{ test.location }}",
    "{{ test.depth_range }}",
    "{{ test.n30 }}",
    "{{ test.lithology }}",
)
MARKER = "Assaigs SPT"


def _set_cell_text(cell, text: str) -> None:
    """Escriu el text conservant el format del primer run de la cel·la."""
    paragraph = cell.paragraphs[0]
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run._element.getparent().remove(run._element)
    else:
        paragraph.add_run(text)


def _find_spt_table(document):
    for table in document.tables:
        first = next((c.text.strip() for row in table.rows for c in row.cells if c.text.strip()), "")
        if first.startswith(MARKER):
            return table
    return None


def main(argv: list[str]) -> int:
    path = Path(argv[0]) if argv else Path("templates/g3dt-jinja-template.docx")
    document = Document(str(path))
    table = _find_spt_table(document)
    if table is None:
        print(f"No s'ha trobat cap taula «{MARKER}» a {path}")
        return 1

    texts = [c.text.strip() for row in table.rows for c in row.cells]
    if any(LOOP_OPEN in t for t in texts):
        print("Ja és un bucle — res a fer.")
        return 0
    if len(table.rows) != 3:
        print(f"Estructura inesperada: {len(table.rows)} files (s'esperaven 3).")
        return 1

    data_tr = table.rows[2]._tr
    open_tr = copy.deepcopy(data_tr)
    close_tr = copy.deepcopy(data_tr)
    data_tr.addprevious(open_tr)
    data_tr.addnext(close_tr)

    for index, marker in ((2, LOOP_OPEN), (4, LOOP_CLOSE)):
        row = table.rows[index]
        _set_cell_text(row.cells[0], marker)
        for cell in row.cells[1:]:
            _set_cell_text(cell, "")
    for cell, text in zip(table.rows[3].cells, DATA_CELLS):
        _set_cell_text(cell, text)

    document.save(str(path))
    print(f"Taula «{MARKER}» convertida en bucle a {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
