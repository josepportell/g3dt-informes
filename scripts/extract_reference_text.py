#!/usr/bin/env python3
"""Extract text from Eva's signed .docx reference reports.

Reads each benchmark .docx, extracts paragraphs and tables, and writes
structured plain-text files for the G3DT benchmark layer.

Usage:
    python scripts/extract_reference_text.py
    python scripts/extract_reference_text.py --benchmarks-dir /other/path
"""

from __future__ import annotations

import argparse
from datetime import datetime, timezone
from pathlib import Path

from docx import Document

_PROJECT_ROOT = Path(__file__).resolve().parent.parent
BENCHMARKS_SRC = Path("/mnt/c/claude/g3dt/benchmarks")

PROJECTS = [
    ("3001621 Castellar del Valles", "3001621_informe_v0.docx"),
    ("3001631 Rubi", "3001631_informe.docx"),
    ("4001607 Linyola", "4001607_informe.docx"),
    ("4001612 Bell-Lloc", "4001612_informe.docx"),
    ("4001670 Alcoletge", "4001670_informe.docx"),
    ("4001671 Vilanova de Segria", "4001671_informe.docx"),
    ("4001679 Anciles", "4001679_informe_V0.docx"),
]


def extract_docx_text(
    docx_path: Path,
) -> tuple[list[str], list[list[list[str]]]]:
    """Extract paragraphs and tables from a .docx file.

    Returns:
        (paragraphs, tables) where tables is a list of tables,
        each table a list of rows, each row a list of cell strings.
    """
    doc = Document(str(docx_path))

    paragraphs = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            paragraphs.append(text)

    tables = []
    for table in doc.tables:
        rows: list[list[str]] = []
        for row in table.rows:
            seen: set[str] = set()
            cells: list[str] = []
            for cell in row.cells:
                text = cell.text.strip()
                cell_id = f"{id(cell._tc)}"
                if cell_id in seen:
                    continue
                seen.add(cell_id)
                cells.append(text)
            rows.append(cells)
        tables.append(rows)

    return paragraphs, tables


def format_output(
    expedient: str,
    municipality: str,
    filename: str,
    paragraphs: list[str],
    tables: list[list[list[str]]],
) -> str:
    """Format extracted data as structured plain text."""
    timestamp = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")

    lines: list[str] = [
        f"# Reference: {expedient} - {municipality}",
        f"# Source: {filename}",
        f"# Extracted: {timestamp}",
        "",
        "=== PARAGRAPHS ===",
    ]

    for para in paragraphs:
        lines.append(para)

    for idx, table in enumerate(tables, 1):
        preview = ""
        if table and table[0]:
            preview = table[0][0][:50]
        lines.append("")
        lines.append(f"=== TABLE {idx}: {preview} ===")
        for row in table:
            lines.append("\t".join(row))

    lines.append("")
    return "\n".join(lines)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Extract text from G3DT benchmark .docx files",
    )
    parser.add_argument(
        "--benchmarks-dir",
        type=Path,
        default=BENCHMARKS_SRC,
        help=f"Source directory with project subfolders (default: {BENCHMARKS_SRC})",
    )
    args = parser.parse_args()

    benchmarks_dir: Path = args.benchmarks_dir
    output_dir = _PROJECT_ROOT / "docs" / "benchmarks"
    output_dir.mkdir(parents=True, exist_ok=True)

    print(f"Source: {benchmarks_dir}")
    print(f"Output: {output_dir}")
    print()

    processed = 0
    errors = 0

    for folder_name, docx_name in PROJECTS:
        parts = folder_name.split(" ", 1)
        expedient = parts[0]
        municipality = parts[1] if len(parts) > 1 else expedient

        docx_path = benchmarks_dir / folder_name / docx_name

        if not docx_path.exists():
            print(f"  SKIP {expedient} ({municipality}): {docx_path} not found")
            errors += 1
            continue

        print(f"  [{processed + 1}/{len(PROJECTS)}] {expedient} {municipality}...", end=" ")

        try:
            paragraphs, tables = extract_docx_text(docx_path)
            text = format_output(expedient, municipality, docx_name, paragraphs, tables)

            out_path = output_dir / f"{expedient}-text.txt"
            out_path.write_text(text, encoding="utf-8")

            print(f"{len(paragraphs)} paragraphs, {len(tables)} tables -> {out_path.name}")
            processed += 1
        except Exception as e:
            print(f"ERROR: {e}")
            errors += 1

    print()
    print(f"Done: {processed} extracted, {errors} errors")


if __name__ == "__main__":
    main()
