#!/usr/bin/env python3
"""
Analyze cover page, watermark/background, and index structure in G3DT samples.
"""

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from pathlib import Path
import xml.etree.ElementTree as ET

def analyze_document(docx_path: str):
    """Analyze document for cover, watermark, index."""
    doc = Document(docx_path)

    print(f"\n{'='*60}")
    print(f"ANALYZING: {Path(docx_path).name}")
    print(f"{'='*60}")

    # 1. First few paragraphs (cover page detection)
    print("\n## FIRST 20 PARAGRAPHS (Cover detection):")
    for i, para in enumerate(doc.paragraphs[:20]):
        text = para.text[:60] + "..." if len(para.text) > 60 else para.text
        style = para.style.name if para.style else "None"
        print(f"  {i:2d}. [{style:20s}] {text}")

    # 2. Sections analysis (multiple sections = different headers/footers)
    print(f"\n## SECTIONS: {len(doc.sections)}")
    for i, section in enumerate(doc.sections):
        print(f"\n  Section {i+1}:")
        print(f"    Page: {section.page_width.cm:.1f} x {section.page_height.cm:.1f} cm")
        print(f"    Margins: T={section.top_margin.cm:.1f}, B={section.bottom_margin.cm:.1f}")
        print(f"    Different first page header: {section.different_first_page_header_footer}")

        # Header content
        if section.header:
            print(f"    Header paragraphs: {len(section.header.paragraphs)}")
            for j, p in enumerate(section.header.paragraphs[:3]):
                print(f"      H{j}: {p.text[:50]}...")

        # Footer content
        if section.footer:
            print(f"    Footer paragraphs: {len(section.footer.paragraphs)}")
            for j, p in enumerate(section.footer.paragraphs[:3]):
                print(f"      F{j}: {p.text[:50]}...")

    # 3. Look for watermark/background in header
    print("\n## WATERMARK/BACKGROUND SEARCH:")

    # Check document relationships for images
    doc_part = doc.part
    print(f"  Document relationships: {len(doc_part.rels)}")

    image_rels = []
    for rel_id, rel in doc_part.rels.items():
        if 'image' in rel.reltype.lower():
            image_rels.append((rel_id, rel.target_ref))
            print(f"    Image: {rel_id} -> {rel.target_ref}")

    # Check header for background shapes
    if doc.sections:
        header = doc.sections[0].header
        header_xml = header._element.xml
        if 'v:background' in header_xml or 'wp:anchor' in header_xml:
            print("  ✓ Found background/anchor elements in header XML")
        if 'w:pict' in header_xml:
            print("  ✓ Found picture elements in header")

    # 4. Index/TOC detection
    print("\n## INDEX/TABLE OF CONTENTS:")
    toc_found = False
    for i, para in enumerate(doc.paragraphs):
        style = para.style.name if para.style else ""
        if 'toc' in style.lower() or 'índex' in para.text.lower() or 'index' in para.text.lower():
            if not toc_found:
                print(f"  TOC starts at paragraph {i}")
                toc_found = True
            if i < 50:  # Show first TOC entries
                print(f"    {i:2d}. [{style:10s}] {para.text[:50]}")

    # 5. Page breaks detection
    print("\n## PAGE BREAKS:")
    page_break_count = 0
    for i, para in enumerate(doc.paragraphs):
        for run in para.runs:
            if run._element.xml.find('w:br') != -1 and 'page' in run._element.xml:
                page_break_count += 1
                print(f"  Page break after paragraph {i}: {para.text[:30]}...")
    print(f"  Total page breaks: {page_break_count}")

    # 6. Look for specific cover elements
    print("\n## COVER PAGE ELEMENTS:")
    cover_styles = ['Title', 'Subtitle', 'Portada', 'Cover']
    for i, para in enumerate(doc.paragraphs[:30]):
        style = para.style.name if para.style else ""
        for cs in cover_styles:
            if cs.lower() in style.lower():
                print(f"  {i}: [{style}] {para.text[:40]}")

if __name__ == "__main__":
    sample = "/home/josep/projects/claudecode-job/clients/g3dt/samples/3001631_informe.docx"
    analyze_document(sample)
