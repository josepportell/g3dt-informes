#!/usr/bin/env python3
"""
Detailed analysis of header structure for cover and watermark.
"""

from docx import Document
from docx.oxml.ns import qn
from pathlib import Path
import re

def analyze_headers(docx_path: str):
    """Analyze header structure in detail."""
    doc = Document(docx_path)

    print(f"\n{'='*60}")
    print("HEADER ANALYSIS")
    print(f"{'='*60}")

    section = doc.sections[0]

    # Check first page header
    print("\n## FIRST PAGE HEADER:")
    if section.first_page_header:
        fph = section.first_page_header
        print(f"  Paragraphs: {len(fph.paragraphs)}")
        for i, p in enumerate(fph.paragraphs):
            print(f"    {i}: style='{p.style.name}', text='{p.text[:50] if p.text else '[empty]'}'")

        # Check for images in first page header
        header_xml = fph._element.xml
        drawing_count = header_xml.count('<w:drawing')
        pict_count = header_xml.count('<w:pict')
        print(f"  Drawing elements: {drawing_count}")
        print(f"  Picture elements: {pict_count}")

        # Look for specific image references
        image_refs = re.findall(r'r:embed="(rId\d+)"', header_xml)
        print(f"  Image references: {image_refs}")
    else:
        print("  No first page header defined")

    # Check regular header
    print("\n## REGULAR HEADER:")
    if section.header:
        h = section.header
        print(f"  Paragraphs: {len(h.paragraphs)}")
        for i, p in enumerate(h.paragraphs):
            print(f"    {i}: style='{p.style.name}', text='{p.text[:50] if p.text else '[empty]'}'")

        header_xml = h._element.xml
        drawing_count = header_xml.count('<w:drawing')
        print(f"  Drawing elements: {drawing_count}")

    # Check footer
    print("\n## FOOTER:")
    if section.footer:
        f = section.footer
        print(f"  Paragraphs: {len(f.paragraphs)}")
        for i, p in enumerate(f.paragraphs):
            print(f"    {i}: style='{p.style.name}', text='{p.text}'")

    # Extract cover info from first page header XML
    print("\n## FIRST PAGE HEADER XML ELEMENTS:")
    if section.first_page_header:
        xml = section.first_page_header._element.xml

        # Look for text boxes (often used for cover layout)
        if '<w:txbxContent>' in xml:
            print("  ✓ Contains text boxes")

        # Look for shapes
        if '<wps:wsp>' in xml or '<v:shape' in xml:
            print("  ✓ Contains shapes")

        # Look for positioned elements
        if 'wp:anchor' in xml:
            print("  ✓ Contains anchored elements")

        if 'wp:positionH' in xml:
            print("  ✓ Contains positioned elements")

if __name__ == "__main__":
    sample = "/home/josep/projects/claudecode-job/clients/g3dt/samples/3001631_informe.docx"
    analyze_headers(sample)
