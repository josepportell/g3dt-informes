#!/usr/bin/env python3
"""
Analyze styles in G3DT sample documents to understand what we need in base template.
"""

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from collections import defaultdict
import json

def analyze_styles(docx_path: str):
    """Extract all styles used in a document."""
    doc = Document(docx_path)

    print(f"\n{'='*60}")
    print(f"ANALYZING: {docx_path}")
    print(f"{'='*60}")

    # 1. Document styles defined
    print("\n## DEFINED STYLES:")
    styles_info = []
    for style in doc.styles:
        if style.type == 1:  # Paragraph style
            info = {
                'name': style.name,
                'type': 'paragraph',
                'base': style.base_style.name if style.base_style else None
            }
            # Try to get font info
            if style.font:
                if style.font.name:
                    info['font'] = style.font.name
                if style.font.size:
                    info['size'] = style.font.size.pt
                if style.font.color and style.font.color.rgb:
                    info['color'] = str(style.font.color.rgb)
            styles_info.append(info)
            print(f"  - {style.name}")

    # 2. Styles actually used in paragraphs
    print("\n## STYLES USED IN PARAGRAPHS:")
    used_styles = defaultdict(int)
    for para in doc.paragraphs:
        style_name = para.style.name if para.style else "None"
        used_styles[style_name] += 1

    for style, count in sorted(used_styles.items(), key=lambda x: -x[1]):
        print(f"  - {style}: {count} times")

    # 3. Tables analysis
    print(f"\n## TABLES: {len(doc.tables)} found")
    for i, table in enumerate(doc.tables):
        rows = len(table.rows)
        cols = len(table.columns)
        print(f"  Table {i+1}: {rows} rows x {cols} cols")
        # Check first cell style
        if table.rows and table.rows[0].cells:
            first_cell = table.rows[0].cells[0]
            if first_cell.paragraphs:
                print(f"    Header style: {first_cell.paragraphs[0].style.name}")

    # 4. Sections (headers/footers)
    print(f"\n## SECTIONS: {len(doc.sections)} found")
    for i, section in enumerate(doc.sections):
        print(f"  Section {i+1}:")
        print(f"    Page size: {section.page_width.cm:.1f} x {section.page_height.cm:.1f} cm")
        print(f"    Margins: T={section.top_margin.cm:.1f}, B={section.bottom_margin.cm:.1f}, L={section.left_margin.cm:.1f}, R={section.right_margin.cm:.1f} cm")
        if section.header:
            print(f"    Has header: Yes ({len(section.header.paragraphs)} paragraphs)")
        if section.footer:
            print(f"    Has footer: Yes ({len(section.footer.paragraphs)} paragraphs)")

    return {
        'styles': styles_info,
        'used_styles': dict(used_styles),
        'tables': len(doc.tables),
        'sections': len(doc.sections)
    }

def main():
    samples = [
        "/home/josep/projects/claudecode-job/clients/g3dt/samples/3001631_informe.docx",
    ]

    for sample in samples:
        try:
            analyze_styles(sample)
        except Exception as e:
            print(f"Error analyzing {sample}: {e}")

if __name__ == "__main__":
    main()
