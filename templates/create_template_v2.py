#!/usr/bin/env python3
"""
Create a Jinja2-enabled template from a G3DT sample report.
VERSION 2: Handles text split across multiple runs.

Usage:
    python create_template_v2.py --input sample.docx --output template.docx
    python create_template_v2.py --dry-run
"""

import argparse
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt


# =============================================================================
# REPLACEMENT MAPPINGS
# =============================================================================

# Text replacements (old_text -> new_text)
REPLACEMENTS = [
    # Client name
    ('SRA. JOANA MARTINEZ', '{{ client }}'),

    # Expedition number
    ('3001631', '{{ expedient }}'),

    # Location variations
    ('RUBÍ', '{{ location }}'),
    ('Rubí', '{{ location }}'),

    # Building data
    ('PB + Porxo', '{{ plantes }}'),
    ('951', '{{ superficie_parcela }}'),
    ('92', '{{ superficie_construida }}'),

    # CTE Classification
    ('C-0', '{{ cte_edificacio }}'),
    ('T-1', '{{ cte_sol }}'),

    # Dates
    ('14 de novembre de 2025', '{{ data_camp_text }}'),
    ('17 de desembre de 2025', '{{ data_signatura_text }}'),
]

# Contexts to skip (paragraphs containing these should not be modified)
SKIP_CONTEXTS = [
    'Real decreto',  # Legal references
    'B.O.E.',        # Official bulletins
    'UNE',           # Standards references
    'Hoja 392',      # Map references
]


def get_paragraph_text(paragraph):
    """Get full text from paragraph by joining all runs."""
    return ''.join(run.text for run in paragraph.runs)


def replace_in_paragraph_runs(paragraph, old_text: str, new_text: str) -> bool:
    """
    Replace text across runs in a paragraph.
    Returns True if replacement was made.
    """
    # Get full paragraph text
    full_text = get_paragraph_text(paragraph)

    if old_text not in full_text:
        return False

    # Check if we should skip this paragraph
    for skip in SKIP_CONTEXTS:
        if skip in full_text:
            return False

    # Strategy: If text spans multiple runs, we need to handle it carefully
    # Simple case: text is in a single run
    for run in paragraph.runs:
        if old_text in run.text:
            run.text = run.text.replace(old_text, new_text)
            return True

    # Complex case: text spans multiple runs
    # We need to find where it starts and ends, then consolidate

    # Find the starting position in the full text
    start_pos = full_text.find(old_text)
    if start_pos == -1:
        return False

    end_pos = start_pos + len(old_text)

    # Find which runs contain the text
    current_pos = 0
    runs_to_modify = []

    for i, run in enumerate(paragraph.runs):
        run_start = current_pos
        run_end = current_pos + len(run.text)

        # Check if this run overlaps with our target
        if run_end > start_pos and run_start < end_pos:
            runs_to_modify.append({
                'index': i,
                'run': run,
                'start': run_start,
                'end': run_end,
                'overlap_start': max(run_start, start_pos),
                'overlap_end': min(run_end, end_pos),
            })

        current_pos = run_end

    if not runs_to_modify:
        return False

    # Modify the runs
    # First run: keep text before the target, add the replacement
    # Middle runs: clear them
    # Last run: keep text after the target

    first_run_info = runs_to_modify[0]
    first_run = first_run_info['run']

    # Calculate what to keep from the first run
    text_before = first_run.text[:start_pos - first_run_info['start']]

    # If there's only one run involved
    if len(runs_to_modify) == 1:
        text_after = first_run.text[end_pos - first_run_info['start']:]
        first_run.text = text_before + new_text + text_after
    else:
        # Multiple runs involved
        last_run_info = runs_to_modify[-1]
        last_run = last_run_info['run']
        text_after = last_run.text[end_pos - last_run_info['start']:]

        # Set first run to text_before + replacement
        first_run.text = text_before + new_text

        # Clear middle runs
        for info in runs_to_modify[1:-1]:
            info['run'].text = ''

        # Set last run to text_after
        last_run.text = text_after

    return True


def replace_in_table_cell(cell, old_text: str, new_text: str) -> bool:
    """Replace text in a table cell."""
    cell_text = cell.text.strip()

    # For table cells, only do exact matches to avoid false positives
    if cell_text == old_text:
        for para in cell.paragraphs:
            for run in para.runs:
                if old_text in run.text:
                    run.text = run.text.replace(old_text, new_text)
                    return True
            # If not found in runs, might be directly in paragraph
            if old_text == para.text.strip():
                # Clear all runs and set first one
                if para.runs:
                    para.runs[0].text = new_text
                    for run in para.runs[1:]:
                        run.text = ''
                    return True

    return False


def process_document(doc: Document, dry_run: bool = False) -> list:
    """Process entire document and make replacements."""
    changes = []

    # Process paragraphs
    for i, para in enumerate(doc.paragraphs):
        for old_text, new_text in REPLACEMENTS:
            if old_text in get_paragraph_text(para):
                # Check skip contexts
                skip = False
                for ctx in SKIP_CONTEXTS:
                    if ctx in get_paragraph_text(para):
                        skip = True
                        break

                if not skip:
                    changes.append({
                        'type': 'paragraph',
                        'index': i,
                        'old': old_text,
                        'new': new_text,
                        'context': para.text[:60],
                    })
                    if not dry_run:
                        replace_in_paragraph_runs(para, old_text, new_text)

    # Process tables
    for ti, table in enumerate(doc.tables):
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                cell_text = cell.text.strip()
                for old_text, new_text in REPLACEMENTS:
                    if cell_text == old_text:
                        changes.append({
                            'type': 'table_cell',
                            'table': ti,
                            'row': ri,
                            'col': ci,
                            'old': old_text,
                            'new': new_text,
                        })
                        if not dry_run:
                            replace_in_table_cell(cell, old_text, new_text)

    # Process headers
    for section in doc.sections:
        if section.header:
            for para in section.header.paragraphs:
                for old_text, new_text in REPLACEMENTS:
                    if old_text in get_paragraph_text(para):
                        changes.append({
                            'type': 'header',
                            'old': old_text,
                            'new': new_text,
                        })
                        if not dry_run:
                            replace_in_paragraph_runs(para, old_text, new_text)

    return changes


def main():
    parser = argparse.ArgumentParser(description='Create Jinja2 template v2')
    parser.add_argument('--input', '-i', default='g3dt-base-template.docx',
                        help='Input document path')
    parser.add_argument('--output', '-o', default='g3dt-jinja-template.docx',
                        help='Output template path')
    parser.add_argument('--dry-run', '-n', action='store_true',
                        help='Show what would be replaced')
    args = parser.parse_args()

    input_path = Path(args.input)
    if not input_path.exists():
        print(f"Error: File not found: {input_path}")
        return 1

    print(f"Processing: {input_path}")
    print(f"Output: {args.output}")
    print(f"Dry run: {args.dry_run}")

    doc = Document(str(input_path))
    changes = process_document(doc, dry_run=args.dry_run)

    # Print summary
    print(f"\n{'=' * 60}")
    print(f"TEMPLATE CREATION SUMMARY (v2)")
    print(f"{'=' * 60}")
    print(f"\nTotal replacements: {len(changes)}")

    for i, change in enumerate(changes, 1):
        if change['type'] == 'paragraph':
            print(f"\n  {i}. [Para {change['index']}]")
        elif change['type'] == 'table_cell':
            print(f"\n  {i}. [Table {change['table']}, R{change['row']}, C{change['col']}]")
        else:
            print(f"\n  {i}. [{change['type']}]")
        print(f"      {change['old']} -> {change['new']}")

    if not args.dry_run:
        doc.save(args.output)
        print(f"\n✅ Template saved to: {args.output}")
    else:
        print(f"\n⚠️  DRY RUN - no changes made")

    return 0


if __name__ == '__main__':
    exit(main())
