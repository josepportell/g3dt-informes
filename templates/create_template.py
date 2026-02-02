#!/usr/bin/env python3
"""
Create a Jinja2-enabled template from a G3DT sample report.
Replaces specific text with Jinja2 placeholders while preserving formatting.

Usage:
    python create_template.py --input sample.docx --output template.docx
    python create_template.py --dry-run  # Show what would be replaced
"""

import argparse
import re
from pathlib import Path
from copy import deepcopy
from docx import Document


# =============================================================================
# REPLACEMENT MAPPINGS
# =============================================================================

# Simple text replacements (exact match or pattern -> placeholder)
# NOTE: Order matters - more specific patterns first
TEXT_REPLACEMENTS = [
    # Client name (appears in multiple places)
    ('SRA. JOANA MARTINEZ', '{{ client }}'),
    ('SRA.JOANA MARTINEZ', '{{ client }}'),

    # Expedition number (only in specific context)
    ('Expedient Núm.:  3001631', 'Expedient Núm.:  {{ expedient }}'),

    # Location (be specific to avoid partial matches)
    ('municipi de RUBÍ', 'municipi de {{ location }}'),
    ('municipi de Rubí', 'municipi de {{ location }}'),
    ('al terme municipal de Rubí', 'al terme municipal de {{ location }}'),

    # Building data (only in table context - these are unique values)
    ('PB + Porxo', '{{ plantes }}'),
]

# Table-specific replacements (for table cells only)
TABLE_REPLACEMENTS = [
    # Building data table (Table 0)
    ('951', '{{ superficie_parcela }}'),
    ('92', '{{ superficie_construida }}'),

    # CTE Classification table (Table 1)
    # Be careful - only replace standalone C-0 and T-1, not part of other text
]

# CTE values need special handling to avoid SPT-1, etc.
CTE_REPLACEMENTS = [
    # These are applied only to Table 1 (CTE classification table)
    ('C-0', '{{ cte_edificacio }}'),
    ('T-1', '{{ cte_sol }}'),
]

# Regex patterns for more complex replacements
REGEX_REPLACEMENTS = [
    # Date pattern: 14 de novembre de 2025
    (r'(\d{1,2})\s+de\s+(gener|febrer|març|abril|maig|juny|juliol|agost|setembre|octubre|novembre|desembre)\s+de\s+(\d{4})',
     '{{ data_camp_text }}'),

    # Date pattern: 17/12/25 or 14/11/2025
    (r'\b(\d{2}/\d{2}/(?:\d{2}|\d{4}))\b', '{{ data }}'),
]


def replace_in_paragraph(paragraph, replacements: list, regex_replacements: list, dry_run: bool = False) -> list:
    """Replace text in a paragraph while preserving formatting."""
    changes = []

    # Get full paragraph text
    full_text = paragraph.text
    if not full_text.strip():
        return changes

    # Check for simple replacements
    for old_text, new_text in replacements:
        if old_text in full_text:
            changes.append({
                'type': 'text',
                'old': old_text,
                'new': new_text,
                'context': full_text[:80],
            })
            if not dry_run:
                # Replace in runs to preserve formatting
                for run in paragraph.runs:
                    if old_text in run.text:
                        run.text = run.text.replace(old_text, new_text)

    # Check for regex replacements
    for pattern, replacement in regex_replacements:
        matches = re.findall(pattern, full_text, re.IGNORECASE)
        if matches:
            changes.append({
                'type': 'regex',
                'pattern': pattern,
                'new': replacement,
                'matches': matches[:3],
                'context': full_text[:80],
            })
            if not dry_run:
                # Replace in runs
                for run in paragraph.runs:
                    run.text = re.sub(pattern, replacement, run.text, flags=re.IGNORECASE)

    return changes


def replace_in_table(table, table_index: int, replacements: list, dry_run: bool = False) -> list:
    """Replace text in table cells with table-specific logic."""
    changes = []

    # Determine which replacements to use based on table index
    table_specific = list(replacements)  # Start with general text replacements

    # Add table-specific replacements
    if table_index == 0:  # Building data table
        table_specific.extend(TABLE_REPLACEMENTS)
    elif table_index == 1:  # CTE classification table
        table_specific.extend(CTE_REPLACEMENTS)

    for row_idx, row in enumerate(table.rows):
        for cell_idx, cell in enumerate(row.cells):
            cell_text = cell.text.strip()
            if not cell_text:
                continue

            for old_text, new_text in table_specific:
                # Only match exact cell content to avoid partial matches like "SPT-1" matching "T-1"
                if old_text == cell_text:
                    changes.append({
                        'type': 'table_cell',
                        'location': f'Table {table_index}, row {row_idx}, col {cell_idx}',
                        'old': old_text,
                        'new': new_text,
                        'context': cell_text[:50],
                    })
                    if not dry_run:
                        # Replace in all paragraphs in cell
                        for para in cell.paragraphs:
                            for run in para.runs:
                                if old_text in run.text:
                                    run.text = run.text.replace(old_text, new_text)

    return changes


def create_template(input_path: str, output_path: str, dry_run: bool = False) -> dict:
    """Create a Jinja2 template from a sample document."""
    doc = Document(input_path)
    all_changes = []

    # Process paragraphs
    for i, para in enumerate(doc.paragraphs):
        changes = replace_in_paragraph(
            para,
            TEXT_REPLACEMENTS,
            REGEX_REPLACEMENTS,
            dry_run=dry_run
        )
        for c in changes:
            c['paragraph_index'] = i
        all_changes.extend(changes)

    # Process tables
    for i, table in enumerate(doc.tables):
        changes = replace_in_table(table, i, TEXT_REPLACEMENTS, dry_run=dry_run)
        all_changes.extend(changes)

    # Process headers
    for section in doc.sections:
        if section.header:
            for para in section.header.paragraphs:
                changes = replace_in_paragraph(
                    para,
                    TEXT_REPLACEMENTS,
                    REGEX_REPLACEMENTS,
                    dry_run=dry_run
                )
                for c in changes:
                    c['location'] = 'header'
                all_changes.extend(changes)

        if section.footer:
            for para in section.footer.paragraphs:
                changes = replace_in_paragraph(
                    para,
                    TEXT_REPLACEMENTS,
                    REGEX_REPLACEMENTS,
                    dry_run=dry_run
                )
                for c in changes:
                    c['location'] = 'footer'
                all_changes.extend(changes)

    # Save if not dry run
    if not dry_run:
        doc.save(output_path)

    return {
        'changes': all_changes,
        'total_changes': len(all_changes),
    }


def print_changes(result: dict):
    """Print summary of changes."""
    print("\n" + "=" * 70)
    print("TEMPLATE CREATION SUMMARY")
    print("=" * 70)

    print(f"\nTotal replacements: {result['total_changes']}")

    if result['changes']:
        print("\nDetailed changes:")
        for i, change in enumerate(result['changes'], 1):
            location = change.get('paragraph_index', change.get('table_index', change.get('location', '?')))
            print(f"\n  {i}. [{location}] {change['type']}")
            if 'old' in change:
                print(f"      OLD: {change['old']}")
            if 'pattern' in change:
                print(f"      PATTERN: {change['pattern']}")
            print(f"      NEW: {change['new']}")
            print(f"      CONTEXT: {change.get('context', '')[:60]}...")


def main():
    parser = argparse.ArgumentParser(description='Create Jinja2 template from sample')
    parser.add_argument('--input', '-i', default='g3dt-base-template.docx',
                        help='Input document path')
    parser.add_argument('--output', '-o', default='g3dt-jinja-template.docx',
                        help='Output template path')
    parser.add_argument('--dry-run', '-n', action='store_true',
                        help='Show what would be replaced without making changes')
    args = parser.parse_args()

    input_path = Path(args.input)
    if not input_path.exists():
        print(f"Error: File not found: {input_path}")
        return 1

    print(f"Processing: {input_path}")
    print(f"Output: {args.output}")
    print(f"Dry run: {args.dry_run}")

    result = create_template(str(input_path), args.output, dry_run=args.dry_run)
    print_changes(result)

    if not args.dry_run:
        print(f"\n✅ Template saved to: {args.output}")
    else:
        print(f"\n⚠️  DRY RUN - no changes made")

    return 0


if __name__ == '__main__':
    exit(main())
